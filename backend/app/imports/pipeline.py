"""Orquestração: upload → validação → prévia → confirmação → carga (com versão e auditoria)."""

import io
import logging
import threading
import time
from collections import Counter
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, insert, select
from sqlalchemy.orm import Session

from app.imports.base import StructureError, load_sheets
from app.imports.compare import compare
from app.imports.detector import detect_and_parse
from app.imports.loaders import LOADERS
from app.imports.resolver import Dimensions, final_status, validate
from app.models import DatasetVersion, ImportBatch, ImportError_, ImportRow
from app.services import audit
from app.services.storage import LocalStorage

log = logging.getLogger(__name__)

ALLOWED_SUFFIXES = (".xlsx", ".xlsm", ".csv")


class ImportStateError(Exception):
    pass


def create_batch(
    db: Session,
    storage: LocalStorage,
    *,
    content: bytes,
    file_name: str,
    user_id: int | None,
    dataset_type: str | None,
    options: dict,
) -> ImportBatch:
    suffix = Path(file_name).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise StructureError("Formato não suportado. Envie .xlsx, .xlsm ou .csv")
    path, digest = storage.save(content, suffix)
    previous = db.scalar(
        select(ImportBatch)
        .where(ImportBatch.file_sha256 == digest, ImportBatch.status == "COMPLETED")
        .order_by(ImportBatch.id.desc())
    )
    batch = ImportBatch(
        dataset_type=dataset_type,
        file_name=file_name,
        file_sha256=digest,
        storage_path=path,
        file_size=len(content),
        status="UPLOADED",
        options=options,
        uploaded_by=user_id,
        summary={"same_file_imported_in": previous.id} if previous else None,
    )
    db.add(batch)
    db.flush()
    audit.record(
        db,
        user_id=user_id,
        action="UPLOAD",
        entity_type="import_batch",
        entity_id=batch.id,
        after={"file_name": file_name, "sha256": digest, "dataset_type": dataset_type, "options": options},
    )
    db.commit()
    return batch


def validate_batch(db: Session, batch: ImportBatch, storage: LocalStorage) -> ImportBatch:
    batch.status = "VALIDATING"
    db.commit()
    options = batch.options or {}
    try:
        sheets = load_sheets(storage.read(batch.storage_path), batch.file_name)
        if batch.dataset_type == "PROJECTION":
            # meses que já têm realizado (KSB1) ficam de fora da projeção, por empresa × ano
            options = dict(options) | {
                "projection_closed": {
                    k.split(":", 1)[1]: int(v or 0)
                    for k, v in db.execute(
                        select(DatasetVersion.scope_key, DatasetVersion.last_closed_period).where(
                            DatasetVersion.dataset_type == "ACTUAL", DatasetVersion.is_current
                        )
                    )
                }
            }
        result = detect_and_parse(sheets, batch.dataset_type, options)
    except StructureError as exc:
        batch.status = "FAILED"
        batch.error_message = str(exc)
        db.add(ImportError_(batch_id=batch.id, code="STRUCTURE", severity="ERROR", message=str(exc)))
        db.commit()
        return batch
    except Exception as exc:  # layout inesperado não pode deixar o lote preso em VALIDATING
        log.exception("Erro inesperado ao ler a importação #%s", batch.id)
        batch.status = "FAILED"
        batch.error_message = f"Erro inesperado ao ler o arquivo: {exc}"
        db.add(ImportError_(batch_id=batch.id, code="STRUCTURE", severity="ERROR", message=batch.error_message))
        db.commit()
        return batch

    validate(result, Dimensions(db), options)
    comparison = compare(db, result, options)
    batch.dataset_type = result.dataset_type
    batch.layout = result.layout

    db.execute(ImportRow.__table__.delete().where(ImportRow.batch_id == batch.id))
    db.execute(ImportError_.__table__.delete().where(ImportError_.batch_id == batch.id))
    counts: Counter[str] = Counter()
    row_payload, error_payload = [], []
    for issue in result.structural:
        error_payload.append(
            {
                "batch_id": batch.id,
                "sheet": None,
                "row_number": None,
                "column": issue.column,
                "code": issue.code,
                "severity": issue.severity,
                "message": issue.message,
                "value": None if issue.value is None else str(issue.value),
            }
        )
    for rec in result.records:
        status = final_status(rec)
        counts[status] += 1
        row_payload.append(
            {
                "batch_id": batch.id,
                "record_type": rec.record_type,
                "sheet": rec.sheet,
                "row_number": rec.row_number,
                "natural_key": rec.natural_key,
                "status": status,
                "data": rec.data,
            }
        )
        if status == "DUPLICATE":
            error_payload.append(
                {
                    "batch_id": batch.id,
                    "sheet": rec.sheet,
                    "row_number": rec.row_number,
                    "column": None,
                    "code": "DUPLICATE",
                    "severity": "ERROR",
                    "message": f"Registro duplicado no arquivo (chave {rec.natural_key})",
                    "value": None,
                }
            )
        for issue in rec.issues:
            error_payload.append(
                {
                    "batch_id": batch.id,
                    "sheet": rec.sheet,
                    "row_number": rec.row_number,
                    "column": issue.column,
                    "code": issue.code,
                    "severity": issue.severity,
                    "message": issue.message,
                    "value": None if issue.value is None else str(issue.value)[:500],
                }
            )
    for start in range(0, len(row_payload), 5000):
        db.execute(insert(ImportRow), row_payload[start : start + 5000])
    for start in range(0, len(error_payload), 5000):
        db.execute(insert(ImportError_), error_payload[start : start + 5000])

    batch.total_rows = len(result.records)
    batch.valid_rows = counts["VALID"] + counts["WARNING"]
    batch.warning_rows = counts["WARNING"]
    batch.error_rows = counts["ERROR"]
    batch.duplicate_rows = counts["DUPLICATE"]
    batch.summary = (batch.summary or {}) | {"meta": result.meta, "comparison": comparison} | _summarize(result)
    blocking = any(i.severity == "ERROR" for i in result.structural)
    batch.status = "FAILED" if blocking or batch.valid_rows == 0 else "VALIDATED"
    if batch.status == "FAILED" and not batch.error_message:
        batch.error_message = "Nenhum registro válido para importar" if not blocking else result.structural[0].message
    batch.validated_at = datetime.utcnow()
    db.commit()
    return batch


def _summarize(result) -> dict:
    by_type = Counter(r.record_type for r in result.records)
    summary: dict = {"record_types": dict(by_type)}
    if result.dataset_type in ("ACTUAL", "REFERENCE_BUDGET", "PROJECTION"):
        total = Decimal("0")
        ccs, accounts = set(), set()
        new_ccs: dict[str, str | None] = {}
        new_accounts: dict[str, str | None] = {}
        for r in result.records:
            if final_status(r) in ("VALID", "WARNING"):
                total += sum((Decimal(v) for v in r.data.get("values", {}).values()), Decimal("0"))
                total += Decimal(r.data.get("amount") or 0)
                ccs.add(r.data.get("cost_center"))
                accounts.add(r.data.get("account"))
                # cadastros que a carga vai criar (opção "cadastrar automaticamente")
                if r.data.get("cost_center") and r.data.get("cost_center_id") is None:
                    new_ccs.setdefault(r.data["cost_center"], r.data.get("cost_center_name"))
                if r.data.get("account") and r.data.get("account_id") is None:
                    new_accounts.setdefault(r.data["account"], r.data.get("account_name"))
        summary |= {"total_amount": str(total), "cost_centers": len(ccs), "accounts": len(accounts)}
        if new_ccs or new_accounts:
            summary["new_cost_centers"] = [{"code": c, "name": n} for c, n in sorted(new_ccs.items())]
            summary["new_accounts"] = [{"code": c, "name": n} for c, n in sorted(new_accounts.items())]
    actions = Counter(r.data.get("_action") for r in result.records if r.data.get("_action"))
    if actions:
        summary["actions"] = dict(actions)
    return summary


def confirmation_blockers(batch: ImportBatch) -> list[str]:
    """Situações que exigem confirmação explícita (force) para evitar cargas redundantes."""
    summary = batch.summary or {}
    blockers = []
    if summary.get("same_file_imported_in"):
        blockers.append(f"Este mesmo arquivo já foi importado (importação #{summary['same_file_imported_in']}).")
    if (summary.get("comparison") or {}).get("no_changes"):
        blockers.append("O arquivo não traz nenhuma alteração em relação aos dados vigentes.")
    return blockers


def confirm_batch(db: Session, batch: ImportBatch, user_id: int | None, *, force: bool = False) -> ImportBatch:
    if batch.status != "VALIDATED":
        raise ImportStateError(f"Importação em status {batch.status} não pode ser confirmada")
    newer = db.scalar(
        select(ImportBatch.id).where(
            ImportBatch.id != batch.id,
            ImportBatch.status == "COMPLETED",
            ImportBatch.completed_at.is_not(None),
            ImportBatch.completed_at > (batch.validated_at or batch.created_at),
        )
    )
    if newer is not None:
        raise ImportStateError(
            f"A prévia desta importação ficou desatualizada: a importação #{newer} foi aplicada depois dela. "
            "Envie o arquivo novamente para gerar uma prévia atual."
        )
    blockers = confirmation_blockers(batch)
    if blockers and not force:
        raise ImportStateError(" ".join(blockers) + " Para importar mesmo assim, confirme novamente marcando 'forçar'.")
    batch.status = "CONFIRMED"
    batch.confirmed_by = user_id
    batch.confirmed_at = datetime.utcnow()
    audit.record(
        db,
        user_id=user_id,
        action="CONFIRM",
        entity_type="import_batch",
        entity_id=batch.id,
        after={"forced": bool(blockers), "blockers": blockers} if blockers else None,
    )
    db.commit()
    return batch


def reject_batch(db: Session, batch: ImportBatch, user_id: int | None, reason: str | None) -> ImportBatch:
    if batch.status not in ("VALIDATED", "FAILED", "UPLOADED"):
        raise ImportStateError(f"Importação em status {batch.status} não pode ser descartada")
    batch.status = "REJECTED"
    audit.record(db, user_id=user_id, action="REJECT", entity_type="import_batch", entity_id=batch.id, reason=reason)
    db.commit()
    return batch


def load_batch(db: Session, batch: ImportBatch) -> ImportBatch:
    if batch.status not in ("CONFIRMED", "PROCESSING"):
        raise ImportStateError(f"Importação em status {batch.status} não pode ser processada")
    batch.status = "PROCESSING"
    db.commit()
    try:
        stats = LOADERS[batch.dataset_type](db, batch, batch.confirmed_by)
        batch.status = "COMPLETED"
        batch.completed_at = datetime.utcnow()
        batch.summary = (batch.summary or {}) | {"load": stats}
        audit.record(
            db, user_id=batch.confirmed_by, action="LOAD", entity_type="import_batch", entity_id=batch.id, after=stats
        )
        db.commit()
    except Exception as exc:  # carga é atômica: nada é gravado em caso de erro
        db.rollback()
        log.exception("Falha ao carregar importação %s", batch.id)
        batch = db.get(ImportBatch, batch.id)
        batch.status = "FAILED"
        batch.error_message = f"Erro na carga (nenhum dado gravado): {exc}"
        db.commit()
    return batch


def process_pending(db: Session, storage: LocalStorage, limit: int = 5) -> int:
    """Processa lotes pendentes. SKIP LOCKED permite vários workers sem concorrência."""
    done = 0
    for _ in range(limit):
        batch = db.scalar(
            select(ImportBatch)
            .where(ImportBatch.status.in_(("UPLOADED", "CONFIRMED")))
            .order_by(ImportBatch.id)
            .with_for_update(skip_locked=True)
            .limit(1)
        )
        if batch is None:
            db.rollback()
            break
        if batch.status == "UPLOADED":
            validate_batch(db, batch, storage)
        else:
            load_batch(db, batch)
        done += 1
    return done


def error_report(db: Session, batch: ImportBatch) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inconsistências"
    header = ["Aba", "Linha", "Coluna", "Severidade", "Código", "Mensagem", "Valor"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    errors = db.scalars(select(ImportError_).where(ImportError_.batch_id == batch.id).order_by(ImportError_.row_number))
    for e in errors:
        ws.append([e.sheet, e.row_number, e.column, e.severity, e.code, e.message, e.value])
    for col, width in zip("ABCDEFG", (18, 8, 22, 11, 22, 70, 30), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    resumo = wb.create_sheet("Resumo")
    for k, v in (
        ("Arquivo", batch.file_name),
        ("Tipo", batch.dataset_type),
        ("Layout", batch.layout),
        ("Registros encontrados", batch.total_rows),
        ("Válidos", batch.valid_rows),
        ("Com avisos", batch.warning_rows),
        ("Com inconsistências", batch.error_rows),
        ("Duplicados", batch.duplicate_rows),
    ):
        resumo.append([k, v])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def errors_by_code(db: Session, batch_id: int) -> list[dict]:
    rows = db.execute(
        select(ImportError_.code, ImportError_.severity, func.count(), func.min(ImportError_.message))
        .where(ImportError_.batch_id == batch_id)
        .group_by(ImportError_.code, ImportError_.severity)
        .order_by(func.count().desc())
    )
    return [{"code": c, "severity": s, "count": n, "example": m} for c, s, n, m in rows]


class ImportWorker(threading.Thread):
    """Worker em thread (mesmo container). Pode rodar isolado: `python -m app.worker`."""

    def __init__(self, session_factory, storage: LocalStorage, poll_seconds: float = 2.0) -> None:
        super().__init__(daemon=True, name="import-worker")
        self.session_factory = session_factory
        self.storage = storage
        self.poll_seconds = poll_seconds
        self._stop_event = threading.Event()

    def stop(self) -> None:
        self._stop_event.set()

    def run(self) -> None:
        while not self._stop_event.is_set():
            try:
                with self.session_factory() as db:
                    processed = process_pending(db, self.storage)
            except Exception:
                log.exception("Erro no worker de importação")
                processed = 0
            if not processed:
                time.sleep(self.poll_seconds)
