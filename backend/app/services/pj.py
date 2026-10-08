"""Contratos PJ: leitura com bonificação do ano, mutações auditadas (sem valores em R$), fotos e exportação.

Acesso: só quem tem `users.can_view_pj` (checado nas rotas por `core.deps.require_pj`). Escopo (`Scope`): com
`users.can_view_all_pj`, todos os contratos; sem ele ("Da área"), só os contratos dos centros de custo das **áreas**
(`departments`) dos CCs da própria pessoa (gestor do CC + escopos atribuídos) — sempre, mesmo com perfil
Administrador/Controladoria, que no Orçamento vê todos os CCs. Contrato sem CC só aparece em "Todos".
Regras puras (CNPJ, situação, tempo de casa, bonificação proporcional) em `app/domain/rules/pj.py`.
"""

import base64
import io
import re
import threading
import time
import unicodedata
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.deps import own_cost_center_ids
from app.domain.rules import pj as R
from app.models import Area, CostCenter, Department, PjContract, PjPhoto, User
from app.services import audit

PJ_ENTITY = "pj_contract"
MONEY_FIELDS = ("monthly_value", "annual_bonus")
EDITABLE = (
    "name",
    "company_name",
    "cnpj",
    "role",
    "cost_center_id",
    "email",
    "phone",
    "monthly_value",
    "annual_bonus",
    "start_date",
    "end_date",
    "notes",
    "photo_blurred",
)
AUDIT_HIDDEN = {"id", "created_at", "updated_at", "created_by", "updated_by"}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^[0-9()+\-.\s]{8,30}$")
MAX_MONEY = Decimal("9999999999999999.99")
MAX_NOTES = 2000
MAX_PHOTO_BYTES = 400 * 1024
PHOTO_MIMES = ("image/jpeg", "image/png", "image/webp")


class NotFound(Exception):
    pass


class Invalid(Exception):
    """Dado inválido (422), mensagem em pt-BR sem valores."""


class Forbidden(Exception):
    """Fora da área da pessoa (403)."""


# ---------------------------------------------------------------- escopo


@dataclass(frozen=True)
class Scope:
    all: bool
    cost_centers: frozenset[int]  # CCs visíveis (quando não é "Todos")

    def allows(self, cost_center_id: int | None) -> bool:
        return self.all or (cost_center_id is not None and cost_center_id in self.cost_centers)


def area_cost_center_ids(db: Session, user: User) -> set[int]:
    """CCs das áreas da própria pessoa: os CCs próprios (gestor + escopos) e todos os CCs das áreas deles."""
    own = own_cost_center_ids(db, user)
    if not own:
        return set()
    depts = set(
        db.scalars(
            select(CostCenter.department_id).where(CostCenter.id.in_(own), CostCenter.department_id.is_not(None))
        )
    )
    ids = set(own)
    if depts:
        ids.update(db.scalars(select(CostCenter.id).where(CostCenter.department_id.in_(depts))))
    return ids


def scope_for(db: Session, user: User) -> Scope:
    """ "Todos" vê tudo; "Da área" vê só os CCs das áreas dos próprios CCs, sempre — mesmo com perfil global."""
    if user.can_view_all_pj:
        return Scope(True, frozenset())
    return Scope(False, frozenset(area_cost_center_ids(db, user)))


def _check_scope(scope: Scope, contract: PjContract) -> None:
    """Quem vê só a área cria e edita só contratos com centro de custo dentro dela."""
    if scope.all:
        return
    if contract.cost_center_id is None:
        raise Invalid("Informe o centro de custo: o seu acesso aos contratos PJ é só aos CCs da sua área")
    if contract.cost_center_id not in scope.cost_centers:
        raise Forbidden("Centro de custo fora da sua área: você só pode lançar contratos PJ dos CCs que enxerga")


def audit_contract_ids(db: Session, user: User) -> list[str] | None:
    """Auditoria de PJ visível: None = toda (com "Todos"); senão, só a dos contratos do escopo (inclui arquivados)."""
    scope = scope_for(db, user)
    if scope.all:
        return None
    rows = db.execute(select(PjContract.id, PjContract.cost_center_id)).all()
    return [str(cid) for cid, cc_id in rows if scope.allows(cc_id)]


# ---------------------------------------------------------------- consulta registrada (LGPD)

# Consulta aos contratos fica na auditoria no máximo uma vez a cada 30 min por pessoa
PJ_VIEW_WINDOW = 30 * 60
_views: dict[int, float] = {}
_views_lock = threading.Lock()


def reset_pj_view_log() -> None:
    with _views_lock:
        _views.clear()


def log_pj_view(db: Session, user: User) -> None:
    now = time.time()
    with _views_lock:
        last = _views.get(user.id)
        if last is not None and now - last < PJ_VIEW_WINDOW:
            return
        _views[user.id] = now
    audit.record(db, user_id=user.id, action="PJ_VIEW", entity_type=PJ_ENTITY, after={"tela": "Contratos PJ"})
    db.commit()


# ---------------------------------------------------------------- leitura


def _today() -> date:
    return date.today()


def _norm(text: str | None) -> str:
    decomposed = unicodedata.normalize("NFD", text or "")
    return "".join(c for c in decomposed if not unicodedata.combining(c)).lower()


def _year(value: int | None, today: date) -> int:
    if value is None:
        return today.year
    if not 2000 <= value <= 2100:
        raise Invalid("Ano inválido")
    return value


def _money_str(value: Decimal | None) -> str | None:
    return None if value is None else f"{value:.2f}"


class _Lookups:
    """Centros de custo com área (departments) e setor (areas) e quais contratos têm foto."""

    def __init__(self, db: Session) -> None:
        depts = {d.id: d.name for d in db.scalars(select(Department))}
        sectors = {a.id: a.name for a in db.scalars(select(Area))}
        self.ccs: dict[int, dict] = {
            cc.id: {
                "label": f"{cc.code} · {cc.name}",
                "department": depts.get(cc.department_id) if cc.department_id else None,
                "sector": sectors.get(cc.area_id) if cc.area_id else None,
            }
            for cc in db.scalars(select(CostCenter))
        }
        self.photos = set(db.scalars(select(PjPhoto.contract_id)))


def _out(c: PjContract, year: int, today: date, lk: _Lookups) -> dict:
    months = R.bonus_months(c.start_date, c.end_date, year)
    until = min(today, c.end_date) if c.end_date else today
    tenure = R.tenure(c.start_date, until)
    cc = lk.ccs.get(c.cost_center_id) if c.cost_center_id else None
    return {
        "id": c.id,
        "name": c.name,
        "company_name": c.company_name,
        "cnpj": c.cnpj,
        "cnpj_formatted": R.format_cnpj(c.cnpj),
        "role": c.role,
        "cost_center_id": c.cost_center_id,
        "cost_center": cc["label"] if cc else None,
        "department": cc["department"] if cc else None,
        "sector": cc["sector"] if cc else None,
        "email": c.email,
        "phone": c.phone,
        "monthly_value": _money_str(c.monthly_value),
        "annual_bonus": _money_str(c.annual_bonus),
        "start_date": c.start_date.isoformat(),
        "end_date": c.end_date.isoformat() if c.end_date else None,
        "status": R.status(c.end_date, today),
        "notes": c.notes,
        "has_photo": c.id in lk.photos,
        "photo_blurred": bool(c.photo_blurred),
        "tenure": {"years": tenure.years, "months": tenure.months},
        "bonus": {
            "year": year,
            "months": months,
            "due": _money_str(R.bonus_due(c.annual_bonus, months)),
        },
    }


def _active_contracts(db: Session, scope: Scope, cost_center_id: int | None, q: str | None) -> list[PjContract]:
    stmt = select(PjContract).where(PjContract.archived_at.is_(None)).order_by(PjContract.name, PjContract.id)
    if cost_center_id is not None:
        stmt = stmt.where(PjContract.cost_center_id == cost_center_id)
    rows = [c for c in db.scalars(stmt) if scope.allows(c.cost_center_id)]
    term = _norm((q or "").strip())
    if term:
        digits = re.sub(r"[\s./-]", "", term).upper()
        rows = [
            c
            for c in rows
            if term in _norm(c.name) or term in _norm(c.company_name) or (len(digits) >= 3 and digits in c.cnpj)
        ]
    return rows


def list_contracts(
    db: Session,
    user: User,
    *,
    year: int | None = None,
    status: str | None = None,
    cost_center_id: int | None = None,
    q: str | None = None,
) -> dict:
    today = _today()
    year = _year(year, today)
    if status not in (None, "", R.ACTIVE, R.ENDED):
        raise Invalid("Situação inválida")
    lk = _Lookups(db)
    items = [_out(c, year, today, lk) for c in _active_contracts(db, scope_for(db, user), cost_center_id, q)]
    if status:
        items = [i for i in items if i["status"] == status]
    items.sort(key=lambda i: (i["status"] != R.ACTIVE, _norm(i["name"])))
    log_pj_view(db, user)
    return {"year": year, "items": items}


def summary(
    db: Session, user: User, *, year: int | None = None, cost_center_id: int | None = None, q: str | None = None
) -> dict:
    """Indicadores do cabeçalho. Mensal e bonificação anual: contratos ativos hoje. Devida no ano: todos os
    contratos com meses no ano de referência (inclui os encerrados durante o ano)."""
    today = _today()
    year = _year(year, today)
    contracts = _active_contracts(db, scope_for(db, user), cost_center_id, q)
    active = [c for c in contracts if R.status(c.end_date, today) == R.ACTIVE]
    due = sum(
        (R.bonus_due(c.annual_bonus, R.bonus_months(c.start_date, c.end_date, year)) for c in contracts),
        Decimal("0.00"),
    )
    return {
        "year": year,
        "active": len(active),
        "ended": len(contracts) - len(active),
        "monthly_total": _money_str(sum((c.monthly_value for c in active), Decimal("0.00"))),
        "annual_bonus_total": _money_str(sum((c.annual_bonus or Decimal("0") for c in active), Decimal("0.00"))),
        "bonus_due_total": _money_str(due),
    }


def get_contract(db: Session, user: User, contract_id: int, year: int | None = None) -> dict:
    today = _today()
    contract = _load(db, scope_for(db, user), contract_id)
    out = _out(contract, _year(year, today), today, _Lookups(db))
    log_pj_view(db, user)
    return out


def options(db: Session, user: User) -> dict:
    """Listas do formulário e dos filtros: centros de custo ativos do escopo, com área e setor."""
    scope = scope_for(db, user)
    lk = _Lookups(db)
    ccs = [
        cc
        for cc in db.scalars(select(CostCenter).where(CostCenter.is_active.is_(True)).order_by(CostCenter.code))
        if scope.allows(cc.id)
    ]
    return {
        "all": scope.all,
        "cost_centers": [
            {
                "id": cc.id,
                "code": cc.code,
                "name": cc.name,
                "department": lk.ccs[cc.id]["department"],
                "sector": lk.ccs[cc.id]["sector"],
            }
            for cc in ccs
        ],
    }


def _load(db: Session, scope: Scope, contract_id: int) -> PjContract:
    contract = db.get(PjContract, contract_id)
    # fora do escopo responde como inexistente (não revela que o contrato existe)
    if contract is None or contract.archived_at is not None or not scope.allows(contract.cost_center_id):
        raise NotFound("Contrato não encontrado")
    return contract


# ---------------------------------------------------------------- validação


def _text(value: Any, max_len: int, label: str, *, required: bool = False) -> str | None:
    text = re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""
    if not text:
        if required:
            raise Invalid(f"Informe {label}")
        return None
    if len(text) > max_len:
        raise Invalid(f"{label[0].upper()}{label[1:]}: no máximo {max_len} caracteres")
    return text


def _money(value: Any, label: str) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation:
        raise Invalid(f"{label} inválido") from None
    if not number.is_finite():
        raise Invalid(f"{label} inválido")
    if number < 0:
        raise Invalid(f"{label} não pode ser negativo")
    if number > MAX_MONEY:
        raise Invalid(f"{label} inválido")
    return number.quantize(R.CENTS)


def _date(value: Any, label: str) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        result = date.fromisoformat(str(value).strip()[:10])
    except ValueError:
        raise Invalid(f"Data {label} inválida") from None
    if not 1950 <= result.year <= 2100:
        raise Invalid(f"Data {label} inválida")
    return result


def _apply(db: Session, c: PjContract, data: dict) -> None:
    unknown = set(data) - set(EDITABLE)
    if unknown:
        raise Invalid(f"Campo não editável: {', '.join(sorted(unknown))}")
    if "name" in data:
        c.name = _text(data["name"], 200, "o nome da pessoa", required=True)
    if "company_name" in data:
        c.company_name = _text(data["company_name"], 200, "a razão social", required=True)
    if "cnpj" in data:
        c.cnpj = R.normalize_cnpj(data["cnpj"])
    if "role" in data:
        c.role = _text(data["role"], 150, "a função")
    if "cost_center_id" in data:
        cc_id = data["cost_center_id"]
        if cc_id not in (None, ""):
            if db.get(CostCenter, int(cc_id)) is None:
                raise Invalid("Centro de custo não encontrado")
            c.cost_center_id = int(cc_id)
        else:
            c.cost_center_id = None
    if "email" in data:
        email = _text(data["email"], 255, "o e-mail")
        if email and not EMAIL_RE.match(email):
            raise Invalid("E-mail inválido")
        c.email = email.lower() if email else None
    if "phone" in data:
        phone = _text(data["phone"], 30, "o telefone")
        if phone and not PHONE_RE.match(phone):
            raise Invalid("Telefone inválido")
        c.phone = phone
    if "monthly_value" in data:
        value = _money(data["monthly_value"], "Valor mensal")
        if value is None:
            raise Invalid("Informe o valor mensal do contrato")
        c.monthly_value = value
    if "annual_bonus" in data:
        c.annual_bonus = _money(data["annual_bonus"], "Valor da bonificação")
    if "start_date" in data:
        c.start_date = _date(data["start_date"], "de admissão")
    if "end_date" in data:
        c.end_date = _date(data["end_date"], "de término")
    if "notes" in data:
        notes = str(data["notes"]).strip() if data["notes"] is not None else ""
        if len(notes) > MAX_NOTES:
            raise Invalid(f"Observação: no máximo {MAX_NOTES} caracteres")
        c.notes = notes or None
    if "photo_blurred" in data:
        c.photo_blurred = bool(data["photo_blurred"])


def _validate(c: PjContract) -> None:
    for attr, message in (
        ("name", "Informe o nome da pessoa"),
        ("company_name", "Informe a razão social"),
        ("cnpj", "Informe o CNPJ"),
        ("monthly_value", "Informe o valor mensal do contrato"),
    ):
        if getattr(c, attr) is None:
            raise Invalid(message)
    R.validate_dates(c.start_date, c.end_date)


# ---------------------------------------------------------------- auditoria (sem valores em R$)


def _audit_row(c: PjContract) -> dict:
    row = audit.snapshot(c, exclude=AUDIT_HIDDEN | set(MONEY_FIELDS))
    for key in MONEY_FIELDS:  # valores não vão para a auditoria: só se estão preenchidos
        row[key] = None if getattr(c, key) is None else "(sigiloso)"
    return row


def _money_changes(before: dict[str, Decimal | None], c: PjContract) -> tuple[dict, dict]:
    """Campos em R$ que mudaram: antes "(sigiloso)", depois "(alterado)" — nunca os valores."""
    b, a = {}, {}
    for key in MONEY_FIELDS:
        old, new = before[key], getattr(c, key)
        if old != new:
            b[key] = None if old is None else "(sigiloso)"
            a[key] = None if new is None else "(alterado)"
    return b, a


# ---------------------------------------------------------------- mutações


def create_contract(db: Session, user: User, payload: dict, ip: str | None = None) -> dict:
    contract = PjContract(photo_blurred=False, created_by=user.id, updated_by=user.id)
    _apply(db, contract, payload)
    _validate(contract)
    _check_scope(scope_for(db, user), contract)
    db.add(contract)
    db.flush()
    audit.record(
        db,
        user_id=user.id,
        action="CREATE",
        entity_type=PJ_ENTITY,
        entity_id=contract.id,
        after=_audit_row(contract),
        ip=ip,
    )
    db.commit()
    return get_contract(db, user, contract.id)


def update_contract(db: Session, user: User, contract_id: int, payload: dict, ip: str | None = None) -> dict:
    scope = scope_for(db, user)
    contract = _load(db, scope, contract_id)
    before = _audit_row(contract)
    money_before = {k: getattr(contract, k) for k in MONEY_FIELDS}
    try:
        _apply(db, contract, payload)
        _validate(contract)
        _check_scope(scope, contract)
    except Exception:
        db.rollback()  # nada fica gravado quando a edição é recusada
        raise
    contract.updated_by = user.id
    before, after = audit.diff(before, _audit_row(contract))
    money_b, money_a = _money_changes(money_before, contract)
    before, after = (before or {}) | money_b, (after or {}) | money_a
    if after:
        audit.record(
            db,
            user_id=user.id,
            action="UPDATE",
            entity_type=PJ_ENTITY,
            entity_id=contract.id,
            before=before,
            after=after,
            ip=ip,
        )
    db.commit()
    return get_contract(db, user, contract.id)


def archive_contract(db: Session, user: User, contract_id: int, ip: str | None = None) -> dict:
    contract = _load(db, scope_for(db, user), contract_id)
    contract.archived_at = datetime.now(UTC)
    contract.updated_by = user.id
    audit.record(
        db,
        user_id=user.id,
        action="ARCHIVE",
        entity_type=PJ_ENTITY,
        entity_id=contract.id,
        after={"name": contract.name},
        ip=ip,
    )
    db.commit()
    return {"id": contract.id}


# ---------------------------------------------------------------- fotos


def photos(db: Session, user: User) -> dict[str, str]:
    """Fotos dos contratos não arquivados do escopo (data URLs)."""
    scope = scope_for(db, user)
    rows = db.execute(
        select(PjPhoto, PjContract.cost_center_id)
        .join(PjContract, PjContract.id == PjPhoto.contract_id)
        .where(PjContract.archived_at.is_(None))
    ).all()
    return {
        str(p.contract_id): f"data:{p.mime};base64,{base64.b64encode(p.content).decode()}"
        for p, cc_id in rows
        if scope.allows(cc_id)
    }


def set_photo(db: Session, user: User, contract_id: int, content: bytes, mime: str, ip: str | None = None) -> dict:
    _load(db, scope_for(db, user), contract_id)
    if mime not in PHOTO_MIMES:
        raise Invalid("Envie a foto em JPEG, PNG ou WebP")
    if not content or len(content) > MAX_PHOTO_BYTES:
        raise Invalid("Foto vazia ou maior que 400 KB")
    photo = db.get(PjPhoto, contract_id)
    if photo is None:
        db.add(PjPhoto(contract_id=contract_id, content=content, mime=mime))
    else:
        photo.content, photo.mime = content, mime
        photo.updated_at = datetime.now(UTC)
    audit.record(db, user_id=user.id, action="PHOTO_SET", entity_type=PJ_ENTITY, entity_id=contract_id, ip=ip)
    db.commit()
    return {"id": contract_id}


def delete_photo(db: Session, user: User, contract_id: int, ip: str | None = None) -> dict:
    _load(db, scope_for(db, user), contract_id)
    photo = db.get(PjPhoto, contract_id)
    if photo is not None:
        db.delete(photo)
        audit.record(db, user_id=user.id, action="PHOTO_DELETE", entity_type=PJ_ENTITY, entity_id=contract_id, ip=ip)
    db.commit()
    return {"id": contract_id}


# ---------------------------------------------------------------- exportação


def export_xlsx(db: Session, user: User, year: int | None, ip: str | None = None) -> bytes:
    today = _today()
    year = _year(year, today)
    lk = _Lookups(db)
    items = [_out(c, year, today, lk) for c in _active_contracts(db, scope_for(db, user), None, None)]
    items.sort(key=lambda i: (i["status"] != R.ACTIVE, _norm(i["name"])))
    wb = Workbook()
    ws = wb.active
    ws.title = "Contratos PJ"
    ws.append(
        [
            "Nome",
            "Razão social",
            "CNPJ",
            "Função",
            "Centro de custo",
            "Área",
            "Setor",
            "E-mail",
            "Telefone",
            "Admissão",
            "Término",
            "Situação",
            "Valor mensal",
            "Bonificação anual",
            f"Meses em {year}",
            f"Bonificação devida em {year}",
            "Observação",
        ]
    )
    for i in items:
        ws.append(
            [
                i["name"],
                i["company_name"],
                i["cnpj_formatted"],
                i["role"],
                i["cost_center"],
                i["department"],
                i["sector"],
                i["email"],
                i["phone"],
                date.fromisoformat(i["start_date"]),
                date.fromisoformat(i["end_date"]) if i["end_date"] else None,
                "Ativo" if i["status"] == R.ACTIVE else "Encerrado",
                Decimal(i["monthly_value"]),
                Decimal(i["annual_bonus"]) if i["annual_bonus"] else None,
                i["bonus"]["months"],
                Decimal(i["bonus"]["due"]),
                i["notes"],
            ]
        )
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for idx in (9, 10):
            row[idx].number_format = "DD/MM/YYYY"
        for idx in (12, 13, 15):
            row[idx].number_format = "#,##0.00"
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 45)
    audit.record(
        db,
        user_id=user.id,
        action="EXPORT",
        entity_type=PJ_ENTITY,
        after={"year": year, "contracts": len(items)},
        ip=ip,
    )
    db.commit()
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
