"""Justificativas do orçamento do ciclo, numa tela só (regra de 08/10/2026: "justificar tudo").

Tudo o que compõe o orçamento precisa de justificativa, porque é o gestor da área quem defende o número:
- OPEX: cada conta orçada (ou com histórico relevante zerado no orçamento) — `AccountJustification`; a justificativa
  escrita nas linhas do template também vale para a conta;
- Pessoal: cada movimentação (contratação, desligamento, transferência, promoção, reajuste) —
  `PersonnelMovement.reason`;
- CAPEX: cada solicitação/projeto — `CapexProject.justification`.

`collect` devolve os itens com o contexto para escrever (referência × proposto, linhas, cargo e salário…); `save`
grava com as mesmas regras de edição das telas do CC, auditoria e, quando faltava, o registro em Apontamentos.
"""

from collections import defaultdict
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.rules.common import money
from app.domain.rules.personnel import ACTION_NOUNS
from app.models import (
    Area,
    BudgetLine,
    BudgetSubmission,
    CapexProject,
    CostCenter,
    Department,
    Employee,
    JobPosition,
    PersonnelMovement,
)
from app.services import capex as capex_svc
from app.services import opex as opex_svc
from app.services import personnel as personnel_svc

ZERO = Decimal("0")
MODULES = ("OPEX", "PERSONNEL", "CAPEX")
MODULE_LABELS = {"OPEX": "OPEX", "PERSONNEL": "Pessoal", "CAPEX": "CAPEX"}
# chave igual à do apontamento correspondente (o registro em Apontamentos usa a mesma)
FINDING_KIND = {
    "OPEX_ACCOUNT": "OPEX_JUSTIFICATION",
    "PERSONNEL_MOVEMENT": "PERSONNEL_JUSTIFICATION",
    "CAPEX_PROJECT": "CAPEX_NO_JUSTIFICATION",
}
MOVE_NOUNS = ACTION_NOUNS | {"HIRE": "contratação", "TRANSFER": "transferência"}
MONTHS = ("JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ")


class JustificationError(ValueError):
    pass


def key_for(sub: BudgetSubmission, entity: str, ident) -> str:
    return f"{sub.id}:{entity}:{ident}:{FINDING_KIND[entity]}"


def parse_key(key: str) -> tuple[int, str, str]:
    try:
        sub_id, entity, ident, _ = key.split(":", 3)
        if entity not in FINDING_KIND:
            raise ValueError
        return int(sub_id), entity, ident
    except ValueError as exc:
        raise JustificationError("Item de justificativa inválido") from exc


def _s(v) -> str | None:
    return None if v is None else str(money(Decimal(v)))


def _who(db: Session, mv: PersonnelMovement) -> str:
    if mv.movement_type == "HIRE":
        return (mv.attributes or {}).get("position_name") or "vaga"
    emp = db.get(Employee, mv.employee_id) if mv.employee_id else None
    return emp.name if emp else f"#{mv.employee_id}"


def _opex_items(db: Session, ctx, sub: BudgetSubmission, cc: dict) -> list[dict]:
    view = opex_svc.account_view(db, ctx, sub)
    lines = defaultdict(list)
    for line in db.scalars(
        select(BudgetLine).where(BudgetLine.submission_id == sub.id).order_by(BudgetLine.total_amount.desc())
    ):
        lines[line.account_id].append(line)
    out = []
    for row in view["accounts"]:
        if not row["needs_justification"]:
            continue
        acc_lines = lines.get(row["account_id"], [])
        details = []
        for line in acc_lines[:6]:
            label = line.description or "(sem descrição)"
            details.append(f"{label}: {_brl(line.total_amount)}")
        if len(acc_lines) > 6:
            details.append(f"+ {len(acc_lines) - 6} linha(s)")
        line_texts = sorted({(ln.justification or "").strip() for ln in acc_lines} - {""})
        text = (row["justification"] or "").strip()
        out.append(
            cc
            | {
                "key": key_for(sub, "OPEX_ACCOUNT", row["account_id"]),
                "submission_id": sub.id,
                "module": "OPEX",
                "kind": "OPEX_ACCOUNT",
                "subject": f"{row['code']} {row['name']}",
                "group": row["package"] or "Sem pacote",
                "base": row["variation_base"],
                "proposed": row["proposed"],
                "flags": row["flags"],
                "details": details,
                "line_texts": line_texts,
                "text": text,
                "justified": bool(text or line_texts),
            }
        )
    return out


def _personnel_items(db: Session, sub: BudgetSubmission, cc: dict) -> list[dict]:
    out = []
    for mv in db.scalars(
        select(PersonnelMovement).where(PersonnelMovement.submission_id == sub.id).order_by(PersonnelMovement.id)
    ):
        if mv.movement_type not in personnel_svc.REASON_REQUIRED:
            continue
        if mv.employee_id:
            emp = db.get(Employee, mv.employee_id)
            if emp is None or emp.cost_center_id != sub.cost_center_id:
                continue  # colaborador mudou de CC na base: a movimentação não vale aqui
        else:
            emp = None
        attrs = mv.attributes or {}
        pending = personnel_svc.month_pending(mv)
        month = "sem mês" if pending or not mv.effective_month else MONTHS[mv.effective_month - 1]
        details = [f"{personnel_svc.MOVE_LABELS.get(mv.movement_type, mv.movement_type)} · {month}"]
        if emp is not None:
            pos = db.get(JobPosition, emp.position_id) if emp.position_id else None
            details.append(f"Cargo atual: {pos.name if pos else '—'} · salário {_brl(emp.base_salary)}")
        if mv.new_salary:
            details.append(f"Novo salário: {_brl(mv.new_salary)}")
        new_pos = db.get(JobPosition, mv.position_id) if mv.position_id else None
        if new_pos is not None or attrs.get("position_name"):
            details.append(f"Cargo: {new_pos.name if new_pos else attrs.get('position_name')}")
        text = (mv.reason or "").strip()
        noun = MOVE_NOUNS.get(mv.movement_type, mv.movement_type.lower())
        out.append(
            cc
            | {
                "key": key_for(sub, "PERSONNEL_MOVEMENT", mv.id),
                "submission_id": sub.id,
                "module": "PERSONNEL",
                "kind": "PERSONNEL_MOVEMENT",
                "subject": f"{noun[0].upper() + noun[1:]} de {_who(db, mv)}",
                "group": noun[0].upper() + noun[1:],
                "base": _s(emp.base_salary) if emp is not None else None,
                "proposed": _s(mv.new_salary) if mv.new_salary else None,
                "flags": [],
                "details": details,
                "line_texts": [],
                "text": text,
                "justified": bool(text),
            }
        )
    return out


def _capex_items(db: Session, sub: BudgetSubmission, cc: dict) -> list[dict]:
    out = []
    for p in db.scalars(select(CapexProject).where(CapexProject.submission_id == sub.id).order_by(CapexProject.code)):
        total = sum((i.total_value for i in p.items), ZERO)
        details = [f"{i.description or 'item'}: {_brl(i.total_value)}" for i in p.items[:6]]
        if p.description and p.description != p.title:
            details.insert(0, p.description)
        text = (p.justification or "").strip()
        out.append(
            cc
            | {
                "key": key_for(sub, "CAPEX_PROJECT", p.id),
                "submission_id": sub.id,
                "module": "CAPEX",
                "kind": "CAPEX_PROJECT",
                "subject": f"{p.code} · {p.title}",
                "group": "Projeto" if p.is_project else "Solicitação",
                "base": None,
                "proposed": _s(total),
                "flags": [],
                "details": details,
                "line_texts": [],
                "text": text,
                "justified": bool(text),
            }
        )
    return out


def _brl(v) -> str:
    n = money(Decimal(v or 0))
    return "R$ " + f"{n:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def collect(db: Session, ctx, scope: set[int] | None, cost_center_id: int | None = None) -> list[dict]:
    """Itens a justificar dos orçamentos do ciclo (CCs visíveis ao usuário; `scope=None` = todos)."""
    stmt = select(BudgetSubmission).where(
        BudgetSubmission.version_id == ctx.version.id, BudgetSubmission.module.in_(MODULES)
    )
    if scope is not None:
        stmt = stmt.where(BudgetSubmission.cost_center_id.in_(scope or {-1}))
    if cost_center_id:
        stmt = stmt.where(BudgetSubmission.cost_center_id == cost_center_id)
    subs = list(db.scalars(stmt))
    ccs = {c.id: c for c in db.scalars(select(CostCenter).where(CostCenter.id.in_({s.cost_center_id for s in subs})))}
    areas = {a.id: a.name for a in db.scalars(select(Area))}
    depts = {d.id: d.name for d in db.scalars(select(Department))}
    items: list[dict] = []
    for sub in sorted(subs, key=lambda s: (ccs[s.cost_center_id].code, MODULES.index(s.module))):
        c = ccs[sub.cost_center_id]
        cc = {
            "cost_center_id": c.id,
            "cost_center": f"{c.code} · {c.name}",
            "department": depts.get(c.department_id),
            "sector": areas.get(c.area_id),
            "status": sub.status,
        }
        if sub.module == "OPEX":
            items += _opex_items(db, ctx, sub, cc)
        elif sub.module == "PERSONNEL":
            items += _personnel_items(db, sub, cc)
        else:
            items += _capex_items(db, sub, cc)
    return items


def save(db: Session, ctx, sub: BudgetSubmission, entity: str, ident: str, text: str, user_id: int) -> tuple:
    """Grava a justificativa do item. Devolve (antes, depois) para a auditoria."""
    text = (text or "").strip()
    if entity == "OPEX_ACCOUNT":
        if sub.module != "OPEX":
            raise JustificationError("Item não pertence ao OPEX")
        out = opex_svc.set_justification(db, sub, int(ident), text, user_id)
        return out["before"], out["after"]
    if entity == "PERSONNEL_MOVEMENT":
        mv = db.get(PersonnelMovement, int(ident))
        if mv is None or mv.submission_id != sub.id:
            raise JustificationError("Movimentação não encontrada")
        before = {"reason": mv.reason}
        mv.reason, mv.updated_by = text or None, user_id
        db.flush()
        return before, {"reason": mv.reason}
    project = db.get(CapexProject, int(ident))
    if project is None or project.submission_id != sub.id:
        raise JustificationError("Solicitação de CAPEX não encontrada")
    before = {"justification": project.justification}
    capex_svc.update_project(db, project, {"justification": text or None}, user_id)
    return before, {"justification": project.justification}


def export_workbook(items: list[dict], target_year: int, ref_year: int) -> bytes:
    """Excel das justificativas (todas as áreas/CCs do filtro): uma aba por módulo + resumo por CC."""
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")
    miss_fill = PatternFill("solid", fgColor="FDECEC")
    summary = wb.active
    summary.title = "Resumo"
    summary.append([f"Justificativas do orçamento {target_year}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(
        ["Regra: tudo o que compõe o orçamento precisa de justificativa (o gestor da área defende o número)."]
    )
    summary.append([])
    summary.append(["Área", "Setor", "Centro de custo", "Módulo", "Itens", "Justificados", "Faltando"])
    counts: dict[tuple, list[int]] = defaultdict(lambda: [0, 0])
    for i in items:
        k = (i["department"] or "", i["sector"] or "", i["cost_center"], MODULE_LABELS[i["module"]])
        counts[k][0] += 1
        counts[k][1] += 1 if i["justified"] else 0
    for k, (n, ok) in sorted(counts.items()):
        summary.append([*k, n, ok, n - ok])
    for cell in summary[4]:
        cell.fill, cell.font = head_fill, head_font
    sheets = {
        "OPEX": ("OPEX", ["Conta", "Pacote", f"Referência {ref_year} (anual.)", f"Orçamento {target_year}"]),
        "PERSONNEL": ("Pessoal", ["Movimentação", "Tipo", "Salário atual", "Novo salário"]),
        "CAPEX": ("CAPEX", ["Solicitação", "Tipo", "—", f"Valor {target_year}"]),
    }
    for module, (title, cols) in sheets.items():
        ws = wb.create_sheet(title)
        headers = ["Área", "Setor", "Centro de custo", *cols, "Detalhes", "Justificativa", "Situação"]
        ws.append(headers)
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
        for i in (x for x in items if x["module"] == module):
            text = i["text"] or " / ".join(i["line_texts"])
            ws.append(
                [
                    i["department"],
                    i["sector"],
                    i["cost_center"],
                    i["subject"],
                    i["group"],
                    float(i["base"]) if i["base"] is not None else None,
                    float(i["proposed"]) if i["proposed"] is not None else None,
                    "\n".join(i["details"]),
                    text,
                    "Justificado" if i["justified"] else "Falta justificar",
                ]
            )
            if not i["justified"]:
                for cell in ws[ws.max_row]:
                    cell.fill = miss_fill
        for col, width in enumerate([16, 18, 34, 42, 22, 16, 16, 50, 60, 16], start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row[5:7]:
                cell.number_format = "#,##0.00"
            for cell in row[7:9]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "D2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
