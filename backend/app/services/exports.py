"""Exportação do orçamento em Excel (openpyxl): resumo, carga SAP (chave × mês), consolidado, variações,
detalhes de OPEX, CAPEX e Pessoal, e situação do fluxo por CC."""

import io
from collections import defaultdict
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.rules.common import MONTH_LABELS, money
from app.domain.workflow import STATUS_LABELS
from app.models import (
    Account,
    BudgetLine,
    BudgetPackage,
    BudgetSubmission,
    BudgetVersion,
    CapexProject,
    CostCenter,
    User,
)
from app.services import capex as capex_svc
from app.services import consolidation as svc
from app.services import personnel as personnel_svc
from app.services.opex import Context

ZERO = Decimal("0")
MONEY = "#,##0.00;[Red]-#,##0.00"
PCT = "0.0%;[Red]-0.0%"
HEADER_FILL = PatternFill("solid", fgColor="222222")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="222222")
TOTAL_FONT = Font(bold=True)
TOTAL_BORDER = Border(top=Side(style="thin", color="222222"))
FLAG_LABELS = {
    "GROWTH_ABOVE": "Crescimento acima do limite",
    "REDUCTION_ABOVE": "Redução acima do limite",
    "NEW_ACCOUNT": "Conta nova",
    "NO_BUDGET": "Sem orçamento",
}


def _f(value) -> float:
    return float(value or 0)


def _table(
    ws: Worksheet,
    start_row: int,
    headers: list[str],
    rows: list[list],
    money_cols: set[int] = frozenset(),
    pct_cols: set[int] = frozenset(),
    widths: dict[int, int] | None = None,
    total: bool = False,
) -> int:
    """Escreve cabeçalho + linhas; colunas monetárias com formato e soma opcional. Devolve a próxima linha livre."""
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(start_row, c, h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, value)
            if c in money_cols:
                cell.number_format = MONEY
            elif c in pct_cols:
                cell.number_format = PCT
    end = start_row + len(rows)
    if total and rows:
        tr = end + 1
        ws.cell(tr, 1, "Total").font = TOTAL_FONT
        for c in money_cols:
            col = get_column_letter(c)
            cell = ws.cell(tr, c, f"=SUBTOTAL(9,{col}{start_row + 1}:{col}{end})")
            cell.number_format, cell.font = MONEY, TOTAL_FONT
        for c in range(1, len(headers) + 1):
            ws.cell(tr, c).border = TOTAL_BORDER
        end = tr
    for c in range(1, len(headers) + 1):
        width = (widths or {}).get(c) or (14 if c in money_cols else max(10, min(42, len(str(headers[c - 1])) + 4)))
        ws.column_dimensions[get_column_letter(c)].width = width
    return end + 2


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], **kw) -> Worksheet:
    ws = wb.create_sheet(title)
    _table(ws, 1, headers, rows, **kw)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    return ws


def budget_workbook(db: Session, ctx: Context, version: BudgetVersion, scope: set[int] | None, user: User) -> bytes:
    vctx = svc.version_context(ctx, version)
    rows = svc.rows_for(db, ctx, version, scope)
    ref = svc.reference(db, vctx, scope)
    variations = svc.account_variations(rows, ref, vctx)
    y = ctx.target_year
    months = list(MONTH_LABELS)
    wb = Workbook()

    # ---------------------------------------------------------------- Resumo
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = (
        f"Orçamento {y} · versão {version.label} ({'congelada' if version.status == 'FROZEN' else 'em elaboração'})"
    )
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} por {user.name}" + (
        " · escopo: centros de custo selecionados" if scope is not None else " · todos os centros de custo"
    )
    ws["A3"] = (
        f"{ctx.ref_year} anualizado = realizado até o último mês fechado × 12 / meses fechados. "
        "Pessoal = salário × reajuste × multiplicador do contrato."
    )
    ws["A4"] = (
        "Sem cronograma = CAPEX com valor total mas sem distribuição mensal (pendência crítica no CC): "
        "conta no orçamento do ano, mas não vai para a carga SAP até ser distribuído."
    )
    by_module: dict[str, Decimal] = defaultdict(lambda: ZERO)
    for r in rows:
        by_module[r.module] += r.total
    summary = []
    for m in svc.MODULES:
        codes = [c for c, n in ref.natures.items() if n in svc.MODULE_NATURES[m]]
        prev = sum((ref.prev.get(c, ZERO) for c in codes), ZERO)
        ann = sum((ref.ref_annualized.get(c, ZERO) for c in codes), ZERO)
        proposed = by_module[m]
        summary.append(
            [
                svc.MODULE_LABELS[m],
                _f(prev),
                _f(ann),
                _f(proposed),
                _f(proposed - ann),
                (_f(proposed / ann - 1) if ann else None),
            ]
        )
    next_row = _table(
        ws,
        5,
        ["Módulo", f"{ctx.prev_year} realizado", f"{ctx.ref_year} anualizado", f"{y} orçado", "Variação", "Var. %"],
        summary,
        money_cols={2, 3, 4, 5},
        pct_cols={6},
        widths={1: 28},
        total=True,
    )
    monthly = [
        [
            svc.MODULE_LABELS[m],
            *[_f(sum((r.values[i] for r in rows if r.module == m), ZERO)) for i in range(12)],
            _f(sum((r.unscheduled for r in rows if r.module == m), ZERO)),
            _f(by_module[m]),
        ]
        for m in svc.MODULES
    ]
    ws.cell(next_row, 1, f"{y} por mês").font = TOTAL_FONT
    _table(
        ws,
        next_row + 1,
        ["Módulo", *months, "Sem cronograma", "Total"],
        monthly,
        money_cols=set(range(2, 16)),
        widths={1: 28},
        total=True,
    )

    # ---------------------------------------------------------------- Carga SAP (chave × mês)
    keyed: dict[tuple, list[Decimal]] = {}
    for r in rows:
        k = (r.company_code, r.branch_code or "", r.cost_center_code, r.account_code)
        keyed.setdefault(k, [ZERO] * 13)  # 12 meses + sem cronograma
        for i in range(12):
            keyed[k][i] += r.values[i]
        keyed[k][12] += r.unscheduled
    sap = [
        [k[0], k[1], k[2], k[3], "-".join(k), *[_f(v) for v in vals], _f(money(sum(vals, ZERO)))]
        for k, vals in sorted(keyed.items())
        if any(vals)
    ]
    _sheet(
        wb,
        "Carga SAP",
        ["Empresa", "Filial", "Centro de custo", "Conta", "Chave", *months, "Sem cronograma", "Total"],
        sap,
        money_cols=set(range(6, 20)),
        widths={5: 34},
        total=True,
    )

    # ---------------------------------------------------------------- Consolidado
    _sheet(
        wb,
        "Consolidado",
        [
            "Módulo",
            "Empresa",
            "Filial",
            "CC",
            "Nome do CC",
            "Conta",
            "Nome da conta",
            "Pacote",
            "Chave",
            *months,
            "Sem cronograma",
            "Total",
        ],
        [
            [
                svc.MODULE_LABELS[r.module],
                r.company_code,
                r.branch_code,
                r.cost_center_code,
                r.cost_center_name,
                r.account_code,
                r.account_name,
                r.package,
                r.key,
                *[_f(v) for v in r.values],
                _f(r.unscheduled),
                _f(r.total),
            ]
            for r in rows
        ],
        money_cols=set(range(10, 24)),
        widths={5: 36, 7: 34, 8: 22, 9: 34},
        total=True,
    )

    # ---------------------------------------------------------------- Variações por conta
    _sheet(
        wb,
        "Variações por conta",
        [
            "Conta",
            "Nome",
            "Módulo",
            "Pacote",
            f"{ctx.prev_year} realizado",
            f"{ctx.ref_year} até o mês",
            f"{ctx.ref_year} anualizado",
            f"{ctx.ref_year} orçado",
            f"{y} orçado",
            "Variação",
            "Var. %",
            "Alertas",
        ],
        [
            [
                v["account"],
                v["name"],
                svc.MODULE_LABELS.get(v["module"], v["module"]),
                v["package"],
                _f(v["prev_actual"]),
                _f(v["ref_actual_ytd"]),
                _f(v["ref_annualized"]),
                _f(v["ref_budget"]),
                _f(v["proposed"]),
                _f(v["variation"]),
                _f(v["variation_pct"]) if v["variation_pct"] else None,
                ", ".join(FLAG_LABELS.get(f, f) for f in v["flags"]),
            ]
            for v in variations
        ],
        money_cols=set(range(5, 11)),
        pct_cols={11},
        widths={2: 36, 4: 22, 12: 32},
        total=True,
    )

    # ---------------------------------------------------------------- detalhes por módulo (linhas da versão)
    ccs = {c.id: c for c in db.scalars(select(CostCenter))}
    accounts = {a.id: a for a in db.scalars(select(Account))}
    packages = {p.id: p.name for p in db.scalars(select(BudgetPackage))}
    subs = list(db.scalars(select(BudgetSubmission).where(BudgetSubmission.version_id == version.id)))
    if scope is not None:
        subs = [s for s in subs if s.cost_center_id in scope]
    by_module_subs = defaultdict(list)
    for s in subs:
        by_module_subs[s.module].append(s)

    opex_rows = []
    for s in by_module_subs["OPEX"]:
        cc = ccs[s.cost_center_id]
        for line in db.scalars(select(BudgetLine).where(BudgetLine.submission_id == s.id).order_by(BudgetLine.id)):
            acc = accounts.get(line.account_id)
            values = {v.month: v.amount for v in line.values}
            opex_rows.append(
                [
                    cc.code,
                    cc.name,
                    packages.get(line.package_id),
                    acc.code if acc else None,
                    acc.name if acc else None,
                    {"GENERIC": "Lançamento", "TRAVEL": "Viagem", "EVENT": "Evento"}.get(
                        line.line_type, line.line_type
                    ),
                    line.description,
                    line.supplier,
                    line.justification,
                    "Template" if (line.attributes or {}).get("source") == "TEMPLATE" else "Sistema",
                    *[_f(values.get(m, 0)) for m in range(1, 13)],
                    _f(line.total_amount),
                ]
            )
    _sheet(
        wb,
        "OPEX (linhas)",
        [
            "CC",
            "Nome do CC",
            "Pacote",
            "Conta",
            "Nome da conta",
            "Tipo",
            "Descrição",
            "Fornecedor",
            "Justificativa",
            "Origem",
            *months,
            "Total",
        ],
        opex_rows,
        money_cols=set(range(11, 24)),
        widths={2: 30, 5: 30, 7: 36, 9: 36},
        total=True,
    )

    capex_rows = []
    for s in by_module_subs["CAPEX"]:
        cc = ccs[s.cost_center_id]
        for p in db.scalars(select(CapexProject).where(CapexProject.submission_id == s.id).order_by(CapexProject.id)):
            for item in p.items:
                acc = accounts.get(item.account_id)
                values = {v.month: v.amount for v in item.values}
                issues = "; ".join(
                    i["message"] for i in capex_svc.item_issues(vctx, item) + capex_svc.project_issues(p)
                )
                capex_rows.append(
                    [
                        cc.code,
                        cc.name,
                        p.code,
                        p.title,
                        "Sim" if p.is_project else "Não",
                        p.project_type_code,
                        item.item_name,
                        acc.code if acc else None,
                        acc.name if acc else None,
                        _f(item.unit_value),
                        _f(item.quantity),
                        _f(item.total_value),
                        p.justification,
                        *[_f(values.get(m, 0)) for m in range(1, 13)],
                        issues,
                    ]
                )
    _sheet(
        wb,
        "CAPEX (itens)",
        [
            "CC",
            "Nome do CC",
            "Código",
            "Solicitação",
            "Projeto?",
            "Tipo do projeto",
            "Item",
            "Conta",
            "Nome da conta",
            "Vlr unit.",
            "Qtd",
            "Vlr total",
            "Justificativa",
            *months,
            "Pendências",
        ],
        capex_rows,
        money_cols={10, 12, *range(14, 26)},
        widths={2: 30, 4: 32, 6: 28, 9: 28, 13: 36, 26: 48},
        total=True,
    )

    personnel_rows = []
    if version.status == "WORKING":
        scenario = personnel_svc.baseline(db, vctx)
        cc_names = {c.id: f"{c.code} · {c.name}" for c in ccs.values()}
        for cc_id, positions in personnel_svc.build_positions(db, vctx, scope).items():
            cc = ccs.get(cc_id)
            for pos in positions:
                out = personnel_svc.position_out(pos, scenario, cc_names)
                mv = out["movement"]
                personnel_rows.append(
                    [
                        cc.code if cc else None,
                        cc.name if cc else None,
                        {"EMPLOYEE": "Colaborador", "HIRE": "Vaga", "TRANSFER_IN": "Transferência recebida"}[
                            out["kind"]
                        ],
                        out["registration"],
                        out["name"],
                        out["position"],
                        out["contract_type_code"],
                        _f(out["base_salary"]),
                        mv["label"] if mv else "Manter",
                        mv["month"] if mv else None,
                        _f(mv["new_salary"]) if mv and mv["new_salary"] else None,
                        mv["reason"] if mv else None,
                        *[_f(v) for v in out["monthly"]],
                        _f(out["annual"]),
                    ]
                )
    ws_p = _sheet(
        wb,
        "Pessoal (quadro)",
        [
            "CC",
            "Nome do CC",
            "Tipo",
            "Matrícula",
            "Nome",
            "Cargo",
            "Contrato",
            "Salário atual",
            "Ação",
            "Mês",
            "Novo salário",
            "Justificativa",
            *months,
            "Custo no ano",
        ],
        personnel_rows,
        money_cols={8, 11, *range(13, 26)},
        widths={2: 30, 5: 30, 6: 28, 12: 32},
        total=True,
    )
    if version.status != "WORKING":
        ws_p["A3"] = "Versão congelada: o custo de pessoal está na aba Consolidado (fotografia do congelamento)."

    status_rows = []
    sub_by_cc = defaultdict(dict)
    for s in subs:
        sub_by_cc[s.cost_center_id][s.module] = s.status
    totals_cc: dict[int, dict[str, Decimal]] = defaultdict(lambda: defaultdict(lambda: ZERO))
    for r in rows:
        totals_cc[r.cost_center_id][r.module] += r.total
    for cc in sorted(ccs.values(), key=lambda c: c.code):
        if scope is not None and cc.id not in scope:
            continue
        if cc.id not in sub_by_cc and cc.id not in totals_cc:
            continue
        st = sub_by_cc.get(cc.id, {})
        t = totals_cc.get(cc.id, {})
        status_rows.append(
            [
                cc.code,
                cc.name,
                cc.manager_name,
                *[STATUS_LABELS.get(st.get(m, "DRAFT"), st.get(m)) for m in svc.MODULES],
                *[_f(t.get(m, ZERO)) for m in svc.MODULES],
                _f(sum((t.get(m, ZERO) for m in svc.MODULES), ZERO)),
            ]
        )
    _sheet(
        wb,
        "Status por CC",
        [
            "CC",
            "Nome do CC",
            "Gestor",
            "Situação OPEX",
            "Situação CAPEX",
            "Situação Pessoal",
            "OPEX",
            "CAPEX",
            "Pessoal",
            "Total",
        ],
        status_rows,
        money_cols={7, 8, 9, 10},
        widths={2: 34, 3: 24, 4: 22, 5: 22, 6: 22},
        total=True,
    )
    for sheet in wb.worksheets[1:]:
        sheet.sheet_view.showGridLines = True
    wb["Resumo"].sheet_properties.tabColor = "FF5A5F"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
