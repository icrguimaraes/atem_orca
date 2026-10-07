"""Gera planilhas sintéticas no mesmo layout dos templates reais (sem dados pessoais)."""

import io
from datetime import datetime

from openpyxl import Workbook


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def opex_template_bd() -> bytes:
    """Aba BD-Novo: tabelas lado a lado a partir da linha 5 (CC em B, Filiais em E, Pacotes em H)."""
    wb = Workbook()
    wb.active.title = "Instruções"
    ws = wb.create_sheet("BD-Novo")
    ws["O4"], ws["R4"] = "VIAGENS", "MKT"
    headers = {
        "B5": "Centro de Custo",
        "C5": "Denominação de centro de custos",
        "E5": "Local de negócios",
        "F5": "Filial",
        "H5": "Descrição",
        "I5": "Conta do Razão",
        "J5": "Agrupamento DRE",
        "K5": "Pacote GMD",
        "L5": "Detalhamento",
        "O5": "Cargos",
        "P5": "Multiplicadores",
    }
    for cell, value in headers.items():
        ws[cell] = value
    ws["B6"], ws["C6"] = "1050101011", "DADOS E PROJ. APLICADOS A CONTROLADORIA"
    ws["B7"], ws["C7"] = 1050101012, "CENTRO NOVO"  # código numérico (Excel)
    for i, (code, name) in enumerate([("0010", "ARAUCARIA"), ("0001", "MANAUS"), ("0099", "NOVA FILIAL")], start=6):
        ws[f"E{i}"], ws[f"F{i}"] = code, name
    accounts = [
        ("Hospedagem", "6010301001", "Despesas", "Viagens", 0),
        ("Telefonia, Links e Internet", "6010301002", "Despesas", "DTI", "Links de Internet"),
        ("Conta Nova Teste", "6010309999", "Despesas", "Pacote Inédito", 0),
        ("Conta Inválida", "ABC", "Despesas", "Viagens", 0),
    ]
    for i, row in enumerate(accounts, start=6):
        for col, value in zip("HIJKL", row, strict=True):
            ws[f"{col}{i}"] = value
    ws["O6"], ws["P6"] = "Anal./Espec./Coord.", 1
    wb.create_sheet("Realizado 2026")["B3"] = "Empresa"
    return _bytes(wb)


def realizado_wide(rows: list[tuple], year: int = 2026, months: int = 8) -> bytes:
    """Aba 'Realizado 2026' (tabela larga, cabeçalhos de mês como data)."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Realizado {year}"
    ws["B2"] = f"ATEM - REALIZADO {year}"
    header = [
        "Empresa",
        "Filial",
        "Nome Filial",
        "Centro de Custos",
        "Denominação do Centro de Custos",
        "Gestor do CC",
        "Conta Razão",
        "Denominação da Conta do Razão",
        "Pacotes Orçamento",
    ]
    header += [datetime(year, m, 1) for m in range(1, months + 1)] + [f"TOTAL REALIZADO {year}"]
    ws.append([])
    ws.append([None, *header])
    for row in rows:
        ws.append([None, *row])
    return _bytes(wb)


def ksb1_csv(lines: list[tuple]) -> bytes:
    header = (
        "Empresa;Centro de custo;Classe de custo;Denominação da classe de custo;Exercício;Período;"
        "Data de lançamento;Nº documento;Valor/moeda ACC;Moeda;Texto;Fornecedor;Nome do fornecedor"
    )
    body = "\n".join(";".join("" if v is None else str(v) for v in line) for line in lines)
    return (header + "\n" + body + "\n").encode("utf-8")


def quadro_funcionarios(rows: list[tuple]) -> bytes:
    """Layout do template Pessoal: resumo no topo, cabeçalho na linha 15."""
    wb = Workbook()
    ws = wb.active
    ws.title = "QUADRO FUNCIONARIOS"
    ws["C8"], ws["C9"], ws["E9"] = "Custo Real", "Premissa", 0.05
    ws["Y13"] = "6010103005"
    ws["B14"], ws["I14"] = "DADOS DOS FUNCIONÁRIOS", "DADOS DA AÇÃO"
    header = [
        "MATRICULA",
        "NOME",
        "CARGO ATUAL",
        "EMPRESA",
        "DIVISÃO",
        "CENTRO DE CUSTO",
        "SALÁRIO MENSAL R$",
        "AÇÃO",
        "MÊS DA AÇÃO",
        "NOVO CARGO",
        "NOVO SALARIO",
        "TIPO DE CONTRATO",
        "AUXILIO CRECHE",
        "VALE TRANSPORTE",
    ]
    for col, value in enumerate(header, start=2):
        ws.cell(15, col, value)
    for r, row in enumerate(rows, start=16):
        for col, value in enumerate(row, start=2):
            ws.cell(r, col, value)
    return _bytes(wb)


def premissas() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "PREMISSAS MACROECONOMICAS"
    ws["B9"] = "Indicadores Macroeconômicos"
    ws.append([])
    ws["B10"], ws["C10"], ws["D10"], ws["E10"], ws["F10"] = (
        "Indicador",
        "Fonte e Informações Adicionais",
        "Atualização",
        2026,
        2027,
    )
    ws["B11"], ws["C11"], ws["D11"], ws["E11"], ws["F11"] = "IPCA", "BACEN", datetime(2026, 9, 1), 0.04, 0.035
    ws["B12"], ws["C12"], ws["E12"] = "PTAX", "Focus", 5.4
    ws["B14"] = "Premissas de Negócio"
    ws["B15"], ws["C15"], ws["D15"], ws["E15"], ws["F15"] = (
        "Indicador",
        "Fonte e Informações Adicionais",
        "Atualização",
        2026,
        2027,
    )
    ws["B16"], ws["C16"], ws["E16"], ws["F16"] = "Volume de Vendas (k M³)", "Comercial", 2900, 3000
    ws["B17"], ws["E17"], ws["F17"] = "Diesel", 2000, 2100
    ws["B18"], ws["C18"], ws["E18"] = "ATEM", "Crescimento da rede", 88
    return _bytes(wb)


def opex_template_filled(budget_cc: str = "1050101011") -> bytes:
    """Template OPEX como volta do gestor: BD-Novo + Realizado 2026 + abas de pacote preenchidas
    (valores como o Excel grava em cache: CHAVE e códigos já calculados)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(opex_template_bd()))
    del wb["Realizado 2026"]
    real = wb.create_sheet("Realizado 2026")
    real.append([])
    real.append([None, "ATEM - REALIZADO 2026"])
    real.append(
        [
            None,
            "Empresa",
            "Filial",
            "Nome Filial",
            "Centro de Custos",
            "Denominação do Centro de Custos",
            "Gestor do CC",
            "Conta Razão",
            "Denominação da Conta do Razão",
            "Pacotes Orçamento",
            *[datetime(2026, m, 1) for m in range(1, 9)],
            "TOTAL REALIZADO 2026",
        ]
    )
    real.append(
        [
            None,
            "1001",
            "0001",
            "MANAUS",
            budget_cc,
            "DADOS E PROJ.",
            "Gestor",
            "6010301002",
            "Telefonia",
            "DTI",
            *[300] * 8,
            2400,
        ]
    )
    months = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]

    viagens = wb.create_sheet("I - Viagens")
    header = [
        "CHAVE",
        "CHAVE",
        "CHAVE",
        "FILIAL",
        "DIVISÃO",
        "DENOMINAÇÃO DO CENTRO DE CUSTO",
        "CENTRO DE CUSTO",
        "OBJETIVO DA VIAGEM",
        "CARGO",
        "IDA (mês)",
        "VOLTA (mês)",
        "PERÍODO (nº de dias)",
        "ORIGEM",
        "DESTINO",
        "TIPO",
        "DESPESAS COM PASSAGENS",
        "DIÁRIA DE VIAGEM",
        "HOSPEDAGEM",
    ]
    for col, text in enumerate(header, start=2):
        viagens.cell(61, col, text)
    viagens.cell(61, 28, "CHAVE")
    for i, m in enumerate(months):
        viagens.cell(61, 29 + i, m)
    viagens.cell(61, 41, 2027)
    keys = [f"1001-0001-{budget_cc}-{acc}" for acc in ("6010301011", "6010301036", "6010301001")]
    trips = [  # valores como o Excel grava (fórmulas já calculadas); "-" = zero
        ("Auditoria SP", "Gerentes", "MAR", "MAR", 3, "AM", "SP", "Nacional", 1800, "-", 0),
        ("Visita base Belém", "Anal./Espec./Coord.", "MAR", None, 4, "AM", "PA", "Nacional", 0, 0, 2400),
    ]
    for r, t in enumerate(trips, start=62):
        row = [*keys, "MANAUS", "0001", "DADOS E PROJ. APLICADOS A CONTROLADORIA", budget_cc, *t]
        for col, value in enumerate(row, start=2):
            viagens.cell(r, col, value)
    viagens.cell(64, 2, "1001-0-0-6010301011")  # linha vazia do template (só fórmulas)
    # consolidador (só conferência): soma por CHAVE e mês de ida
    viagens.cell(62, 28, keys[0])
    viagens.cell(62, 31, 1800)
    viagens.cell(63, 28, keys[2])
    viagens.cell(63, 31, 2400)
    viagens.cell(64, 28, "1001-0-0-6010301036")

    st = wb.create_sheet("II - Serviços de Terceiros")
    for col, text in enumerate(
        [
            "CHAVE",
            "DETALHAMENTO DO CONTRATO",
            "GESTOR DO CONTRATO",
            "FORNECEDOR",
            "OBSERVAÇÕES",
            "FILIAL",
            "DIVISÃO",
            "DENOMINAÇÃO DO CENTRO DE CUSTOS",
            "CENTRO DE CUSTO",
            "DESCRIÇÃO DA CONTA CONTÁBIL",
            "CONTA CONTÁBIL",
            *months,
            2027,
        ],
        start=2,
    ):
        st.cell(14, col, text)
    st.append(
        [
            None,
            f"1001-0001-{budget_cc}-6010201003",
            "Consultoria tributária",
            "Ana",
            "KPMG",
            "Reajuste IPCA",
            "MANAUS",
            "0001",
            "DADOS",
            budget_cc,
            "Serviços de Consultoria e Assessoria",
            "6010201003",
            *[1000] * 12,
            12000,
        ]
    )
    st.append([None, "1001---", None, None, None, None, None, None, None, None, None, None, *[None] * 12])

    dti = wb.create_sheet("VI - DTI")
    for col, text in enumerate(
        [
            "CHAVE",
            "DETALHAMENTO DA DESPESA",
            "VIGÊNCIA",
            "TIPO DE CONTRATO",
            "FILIAL",
            "DIVISÃO",
            "DENOMINAÇÃO DO CENTRO DE CUSTOS",
            "CENTRO DE CUSTO",
            "PRODUTO/SERVIÇO",
            "DESCRIÇÃO DA CONTA CONTÁBIL",
            "CONTA CONTÁBIL",
            *months,
            2027,
        ],
        start=2,
    ):
        dti.cell(13, col, text)
    dti.append(
        [
            None,
            f"1001-0001-{budget_cc}-6010301002",
            "Links",
            "12m",
            "Manutenção",
            "MANAUS",
            "0001",
            "DADOS",
            budget_cc,
            "Links de Internet",
            "Telefonia",
            "6010301002",
            *[350] * 12,
            4200,
        ]
    )
    return _bytes(wb)


def capex_template_filled(cc: str = "1050101011") -> bytes:
    """Template CAPEX como volta do gestor: BD-Novo (CC, filiais, contas de ativo), LISTA ATIVOS e
    Template_Orç 2027 com cabeçalho na linha 6 e meses datados do ano anterior (como no arquivo real)."""
    wb = Workbook()
    wb.active.title = "Instruções"
    ws = wb.create_sheet("Template_Orç 2027")
    header = [
        "NOME FILIAL",
        "FILIAL",
        "NOME CENTRO DE CUSTO",
        "CENTRO DE CUSTO",
        "DESCRIÇÃO DA CONTA",
        "CONTA",
        "Projeto?",
        "TIPO DO PROJETO",
        "ITEM",
        "DESCRIÇÃO DETALHADA DO ITEM OU PROJETO",
        "VLR UNIT",
        "QTD",
        "VLR TOTAL",
        "JUSTIFICATIVA",
        *[datetime(2026, m, 1) for m in range(1, 13)],
        "Orçamento 2026",
        None,
        "Check",
    ]
    ws.cell(5, 10, "Não preencher")
    for col, value in enumerate(header, start=2):
        ws.cell(6, col, value)

    def row(r, values, months):
        for col, value in enumerate(values, start=2):
            ws.cell(r, col, value)
        for m, amount in months.items():
            ws.cell(r, 15 + m, amount)

    base = ["MANAUS", "0001", "DADOS E PROJ.", cc]
    # projeto com 2 itens (mesmo tipo e justificativa) → 1 solicitação
    just = "Automatizar a conciliação; redução estimada de R$ 50 mil/ano"
    row(
        7,
        [
            *base,
            "Equipamentos de Informática",
            "1020601005",
            "Sim",
            "Automação e Transformação Digital",
            "NOTEBOOK",
            "Notebooks para o time de dados",
            6000,
            3,
            18000,
            just,
        ],
        {3: 18000},
    )
    row(
        8,
        [
            *base,
            "Licenças e Software",
            "1020701002",
            "Sim",
            "Automação e Transformação Digital",
            "LICENÇA",
            "Licença da ferramenta de BI",
            12000,
            1,
            12000,
            just,
        ],
        {3: 6000, 9: 6000},
    )
    # aquisição avulsa com cronograma divergente e valor baixo → avisos
    row(
        9,
        [
            *base,
            "Móveis, Utensílios e Instalações",
            "1020601004",
            "Não",
            None,
            "CADEIRA",
            "Cadeiras ergonômicas",
            1000,
            4,
            4000,
            "Reposição",
        ],
        {5: 3000},
    )
    # conta de despesa → erro
    row(10, [*base, "Telefonia", "6010301002", "Não", None, "LINK", "Link", 2000, 1, 2000, None], {1: 2000})
    row(11, [None] * 12 + [0], {})  # linha do modelo só com fórmula zerada

    bd = wb.create_sheet("BD-Novo")
    for cell, value in {
        "B2": "Centro de Custo",
        "C2": "Denominação de centro de custos",
        "D2": "Gestor do Centro de Custo",
        "F2": "Local de negócios",
        "G2": "Filial",
        "I2": "Descrição",
        "J2": "Conta do Razão",
        "K2": "Agrupamento DRE",
        "L2": "Pacote GMD",
    }.items():
        bd[cell] = value
    bd["B3"], bd["C3"], bd["D3"] = cc, "DADOS E PROJ. APLICADOS A CONTROLADORIA", "Gestor Teste"
    bd["F3"], bd["G3"] = "0001", "MANAUS"
    for i, (name, code) in enumerate(
        [
            ("Equipamentos de Informática", 1020601005),
            ("Licenças e Software", 1020701002),
            ("Móveis, Utensílios e Instalações", 1020601004),
        ],
        start=3,
    ):
        bd[f"I{i}"], bd[f"J{i}"], bd[f"K{i}"], bd[f"L{i}"] = name, code, "Sem Agrupamento", "Capex"

    lista = wb.create_sheet("LISTA ATIVOS")
    lista.append(["Item Principal", "Nome Classe", "ATEM", "REAM", None, None, None, "Item Principal", "Nome Classe"])
    lista.append(
        [
            "NOTEBOOK",
            "Equipamentos de Informática",
            "1020601005",
            "1020601005",
            None,
            None,
            None,
            "NOTEBOOK",
            "Equipamentos de Informática",
        ]
    )
    lista.append(
        ["CADEIRA", "Móveis, Utensílios e Instalações", "1020601004", "1020601004", None, None, None, "#REF!", "#REF!"]
    )
    lista.append(["DRONE", "Classe Nova", "1020601003", None])
    return _bytes(wb)


KSB1_EXPORT_HEADER = [
    "Empresa",
    "Centro",
    "Centro custo",
    "Denominação objeto",
    "Classe de custo",
    "Denom.classe custo",
    "Denominação",
    "Documento de compras",
    "Texto do pedido",
    "Material",
    "Texto breve material",
    "Nº doc.de referência",
    "Tipo de conta de contraparti",
    "Conta lnçto.contrap.",
    "Nº ref.estorno",
    "Denom.conta de contrapartida",
    "Nome do usuário",
    "Valor/moeda objeto",
    "Data do documento",
    "Data de lançamento",
    "Qtd.total entrada",
    "Unid.medida lançada",
    "Data de entrada",
    "Hora do registro",
]


def ksb1_export_xlsx(rows: list[dict], subtotals: dict[str, float] | None = None, total: float | None = None) -> bytes:
    """Exportação SAP KSB1 em xlsx (aba "Data", cabeçalhos abreviados, planta em "Centro", duas datas,
    subtotais por centro de custo e total geral só com o valor preenchido)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(KSB1_EXPORT_HEADER)
    for row in rows:
        ws.append(
            [
                row.get("company", "1001"),
                row.get("plant", "C001"),
                row["cost_center"],
                row.get("cost_center_name", "CENTRO"),
                row["account"],
                row.get("account_name", "CONTA"),
                row.get("text", "Lançamento"),
                None,
                None,
                None,
                None,
                row.get("document", "5100001"),
                "S",
                "2010501001",
                None,
                "Contrapartida",
                "USUARIO",
                row["amount"],
                row.get("document_date"),
                row.get("posting_date"),
                0,
                None,
                row.get("posting_date"),
                "10:00:00",
            ]
        )
    for cc, amount in (subtotals or {}).items():
        ws.append([None, None, cc] + [None] * 14 + [amount] + [None] * 6)
    if total is not None:
        ws.append([None] * 17 + [total] + [None] * 6)
    return _bytes(wb)


def opex_template_ream() -> bytes:
    """Template OPEX da REAM (layout antigo): abas de pacote sem a coluna CHAVE, CC alfanumérico, empresa só pela
    divisão (2001) e aba BD com várias tabelas empilhadas (não é o BD-Novo da ATEM)."""
    wb = Workbook()
    wb.active.title = "Instruções"
    bd = wb.create_sheet("BD")
    for col, value in {"M": "Divisão", "N": "Nome Divisão", "Q": "Cod.Centro", "R": "FILIAL"}.items():
        bd[f"{col}10"] = value
    bd["M11"], bd["N11"], bd["Q11"] = "2001", "Manaus", "C201"
    bd["M40"], bd["N40"] = "Fulano de Tal", "Fulano de Tal"  # outra tabela embaixo, na mesma coluna
    months = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
    header = [
        "FORNECEDOR",
        "OBSERVAÇÕES",
        "FILIAL",
        "DIVISÃO",
        "DENOMINAÇÃO DO CENTRO DE CUSTO",
        "TIPO DE CENTRO DE CUSTO ",
        "CENTRO DE CUSTO",
        "DESCRIÇÃO DA CONTA CONTÁBIL",
        "CONTA CONTÁBIL",
        *months,
        2027,
    ]
    ident = ["Manaus", "2001", "Custos", "W", "RFM6003000"]
    viagens = wb.create_sheet("I - Viagens")
    viagens["B9"], viagens["AF9"], viagens["AH9"] = "AM", "Internacional", 1000  # só premissas
    servicos = wb.create_sheet("II - Serviços de Terceiros")
    servicos.append([])
    for _ in range(12):
        servicos.append([])
    servicos.append([None, "CONTRATO", *header])  # linha 14, sem valores
    consumo = wb.create_sheet("VII - Consumo e Expediente")
    for _ in range(13):
        consumo.append([])
    consumo.append([None, "MATERIAL", *header])
    consumo.append([None, None, None, None, *ident, "W-Alimentação", "6010301008", *[600] * 12, 7200])
    consumo.append([None, None, None, None, *ident, "W-Material de Consumo", "6010301006", *[200] * 12, 2400])
    logistica = wb.create_sheet("IX - Logística e Transporte")
    for _ in range(15):
        logistica.append([])
    logistica.append([None, None, "MATERIAL", *header])  # uma coluna a mais à esquerda
    logistica.append([None, None, None, None, None, *ident, "W-Conduções e Táxis", "6010301025", *[900] * 12, 10800])
    return _bytes(wb)
