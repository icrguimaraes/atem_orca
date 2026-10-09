"""Contratos PJ: acesso só com "Vê contratos PJ", escopo "Da área" por centro de custo, validação (nada obrigatório,
pendências em `missing`), CRUD com auditoria sem valores, fotos, indicadores e as migrações 0009/0011 (up/down sem
drift)."""

import io
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import inspect, text

from app.db import SessionLocal, engine
from app.domain.rules.pj import cnpj_check_digits
from app.models import Base, User
from app.services.pj import area_cost_center_ids, audit_contract_ids, reset_pj_view_log
from tests.conftest import login

TODAY = date(2026, 10, 8)
JPEG = bytes([0xFF, 0xD8, 0xFF, 0xE0]) + b"foto sintetica"
USERS = "/api/v1/users"


def fake_cnpj(n: int) -> str:
    base = f"{(n * 11111111) % 100000000:08d}0001"  # n=3 → 33.333.333/0001-xx
    digits = base + cnpj_check_digits(base)
    return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr("app.services.pj._today", lambda: TODAY)
    reset_pj_view_log()


def create_user(client, admin, email: str, roles: list[str], **extra) -> int:
    resp = client.post(
        USERS,
        headers=admin,
        json={"email": email, "name": email.split("@")[0], "password": "Senha@123", "roles": roles} | extra,
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def grant(client, admin, user_id: int, value: bool = True, all_: bool = False) -> dict:
    """Vê contratos PJ: Não (value=False), Da área (value=True) ou Todos (all_=True)."""
    resp = client.patch(f"{USERS}/{user_id}", headers=admin, json={"can_view_pj": value, "can_view_all_pj": all_})
    assert resp.status_code == 200, resp.text
    return resp.json()


@pytest.fixture
def pj_admin(client, admin):
    me = client.get("/api/v1/auth/me", headers=admin).json()
    grant(client, admin, me["id"], all_=True)
    return admin


def make_structure(client, admin) -> dict:
    """Duas áreas: "Área PJ A" com os CCs A1 e A2 e "Área PJ B" com o CC B1. Devolve ids por código."""
    company = next(c for c in client.get("/api/v1/companies", headers=admin).json() if c["code"] == "1001")["id"]
    depts = {}
    for name in ("Área PJ A", "Área PJ B"):
        resp = client.post("/api/v1/departments", headers=admin, json={"name": name})
        assert resp.status_code == 201, resp.text
        depts[name] = resp.json()["id"]
    ccs = {}
    for code, dept in (("9001", "Área PJ A"), ("9002", "Área PJ A"), ("9003", "Área PJ B")):
        body = {"company_id": company, "code": code, "name": f"CC PJ {code}", "department_id": depts[dept]}
        resp = client.post("/api/v1/cost-centers", headers=admin, json=body)
        assert resp.status_code == 201, resp.text
        ccs[code] = resp.json()["id"]
    return {"departments": depts, "ccs": ccs}


def contract(n: int = 1, **extra) -> dict:
    body = {
        "name": f"PESSOA PJ {n}",
        "company_name": f"EMPRESA FICTICIA {n} LTDA",
        "cnpj": fake_cnpj(n),
        "role": "Consultor",
        "monthly_value": "10000.00",
        "annual_bonus": "12000.00",
        "start_date": "2020-06-20",
    }
    return body | extra


# ---------------------------------------------------------------- acesso


def test_every_endpoint_requires_the_flag(client, admin):
    me = client.get("/api/v1/auth/me", headers=admin).json()
    assert me["can_view_pj"] is False and me["pj_access"] == "NONE"
    endpoints = [
        ("get", "/api/v1/pj", None),
        ("get", "/api/v1/pj/summary", None),
        ("get", "/api/v1/pj/options", None),
        ("get", "/api/v1/pj/photos", None),
        ("get", "/api/v1/pj/export.xlsx", None),
        ("get", "/api/v1/pj/1", None),
        ("post", "/api/v1/pj", contract()),
        ("patch", "/api/v1/pj/1", {"role": "x"}),
        ("delete", "/api/v1/pj/1", None),
        ("delete", "/api/v1/pj/1/photo", None),
    ]
    # administrador sem a permissão: 403 em tudo (a permissão não vem do perfil)
    for method, url, body in endpoints:
        resp = getattr(client, method)(url, headers=admin, **({"json": body} if body else {}))
        assert resp.status_code == 403, (method, url, resp.status_code)
        assert resp.json()["detail"] == "Sem acesso aos contratos PJ"
    put = client.put("/api/v1/pj/1/photo", headers=admin, files={"file": ("f.jpg", JPEG, "image/jpeg")})
    assert put.status_code == 403
    assert client.get("/api/v1/pj").status_code == 401

    # o administrador concede a si mesmo; daí em diante enxerga
    assert grant(client, admin, me["id"])["can_view_pj"] is True
    assert client.get("/api/v1/auth/me", headers=admin).json()["pj_access"] == "AREA"
    assert client.get("/api/v1/pj", headers=admin).status_code == 200

    # gestor: sem a permissão não vê; com ela, vê (independe do perfil); não pode conceder
    manager_id = create_user(client, admin, "gestor@test.com", ["MANAGER"])
    manager = login(client, "gestor@test.com", "Senha@123")
    assert client.get("/api/v1/pj", headers=manager).status_code == 403
    assert client.patch(f"{USERS}/{manager_id}", headers=manager, json={"can_view_pj": True}).status_code == 403
    grant(client, admin, manager_id)
    assert client.get("/api/v1/pj", headers=manager).status_code == 200
    grant(client, admin, manager_id, False)
    assert client.get("/api/v1/pj", headers=manager).status_code == 403

    logs = client.get("/api/v1/audit-logs", headers=admin, params={"entity_type": "user", "action": "UPDATE"}).json()
    changes = [i for i in logs["items"] if "can_view_pj" in (i["after"] or {})]
    assert len(changes) == 3
    assert changes[0]["before"] == {"can_view_pj": True} and changes[0]["after"] == {"can_view_pj": False}


def test_access_levels_via_users_api(client, admin):
    manager_id = create_user(client, admin, "gestor@test.com", ["MANAGER"])
    manager = login(client, "gestor@test.com", "Senha@123")
    assert client.get("/api/v1/auth/me", headers=manager).json()["pj_access"] == "NONE"
    grant(client, admin, manager_id)
    me = client.get("/api/v1/auth/me", headers=manager).json()
    assert (me["pj_access"], me["can_view_pj"], me["can_view_all_pj"]) == ("AREA", True, False)
    grant(client, admin, manager_id, all_=True)
    assert client.get("/api/v1/auth/me", headers=manager).json()["pj_access"] == "ALL"
    listed = next(u for u in client.get(USERS, headers=admin).json() if u["id"] == manager_id)
    assert listed["pj_access"] == "ALL"
    # "Não" limpa também o "Todos"
    resp = client.patch(f"{USERS}/{manager_id}", headers=admin, json={"can_view_pj": False})
    assert (resp.json()["pj_access"], resp.json()["can_view_all_pj"]) == ("NONE", False)
    # "Todos" sem o acesso ao módulo não dá acesso
    client.patch(f"{USERS}/{manager_id}", headers=admin, json={"can_view_all_pj": True})
    assert client.get("/api/v1/auth/me", headers=manager).json()["pj_access"] == "NONE"
    assert client.get("/api/v1/pj", headers=manager).status_code == 403
    logs = client.get(
        "/api/v1/audit-logs", headers=admin, params={"entity_type": "user", "entity_id": manager_id, "action": "UPDATE"}
    ).json()["items"]
    levels = [(log["before"], log["after"]) for log in reversed(logs) if "can_view_all_pj" in (log["after"] or {})]
    assert levels[0] == ({"can_view_all_pj": False}, {"can_view_all_pj": True})
    assert levels[1] == (
        {"can_view_pj": True, "can_view_all_pj": True},
        {"can_view_pj": False, "can_view_all_pj": False},
    )
    # já na criação: Da área / Todos
    area_id = create_user(client, admin, "area@test.com", ["MANAGER"], can_view_pj=True)
    all_id = create_user(client, admin, "todos@test.com", ["CONTROLLER"], can_view_pj=True, can_view_all_pj=True)
    only_all = create_user(client, admin, "so-todos@test.com", ["MANAGER"], can_view_all_pj=True)
    by_id = {u["id"]: u["pj_access"] for u in client.get(USERS, headers=admin).json()}
    assert (by_id[area_id], by_id[all_id], by_id[only_all]) == ("AREA", "ALL", "NONE")


# ---------------------------------------------------------------- validação e CRUD


def test_validation_messages(client, pj_admin):
    def post(**extra):
        return client.post("/api/v1/pj", headers=pj_admin, json=contract(**extra))

    # formato continua validado quando há valor
    assert post(cnpj="11.222.333/0001-82").json()["detail"] == "CNPJ inválido"
    assert post(cnpj="123").json()["detail"] == "CNPJ inválido"
    assert post(monthly_value="-1").json()["detail"] == "Valor mensal não pode ser negativo"
    assert post(annual_bonus="-0.01").json()["detail"] == "Valor da bonificação não pode ser negativo"
    assert post(monthly_value="abc").json()["detail"] == "Valor mensal inválido"
    assert post(end_date="2020-06-19").json()["detail"] == "A data de término não pode ser anterior à admissão"
    assert post(start_date="2026-02-30").json()["detail"] == "Data de admissão inválida"
    assert post(email="sem-arroba").json()["detail"] == "E-mail inválido"
    assert post(phone="abc").json()["detail"] == "Telefone inválido"
    assert post(name="x" * 201).json()["detail"] == "O nome da pessoa: no máximo 200 caracteres"
    assert post(cost_center_id=999).json()["detail"] == "Centro de custo não encontrado"
    assert client.get("/api/v1/pj", headers=pj_admin).json()["items"] == []
    # CNPJ alfanumérico aceito; CNPJ repetido não é bloqueado (decisão)
    assert post(cnpj="12.ABC.345/01DE-35").status_code == 201
    assert post(cnpj="12abc34501de35").status_code == 201
    assert [i["cnpj_formatted"] for i in client.get("/api/v1/pj", headers=pj_admin).json()["items"]] == [
        "12.ABC.345/01DE-35",
        "12.ABC.345/01DE-35",
    ]


def test_nothing_is_required_and_missing_fields_are_flagged(client, pj_admin):
    st = make_structure(client, pj_admin)
    # sem nada: salva e devolve todas as pendências
    empty = client.post("/api/v1/pj", headers=pj_admin, json={})
    assert empty.status_code == 201, empty.text
    item = empty.json()
    assert item["missing"] == [
        "Nome da pessoa",
        "Razão social",
        "CNPJ",
        "Função",
        "Centro de custo",
        "Valor mensal",
        "Data de admissão",
    ]
    assert (item["name"], item["cnpj"], item["cnpj_formatted"], item["monthly_value"]) == (None, None, None, None)
    assert item["start_date"] is None and item["tenure"] is None and item["status"] == "ACTIVE"
    assert item["bonus"] == {"year": 2026, "months": 0, "due": "0.00"}

    # só o nome; vazios explícitos ("", "  ", null) também passam
    only_name = client.post(
        "/api/v1/pj",
        headers=pj_admin,
        json={"name": "SO O NOME", "company_name": "  ", "cnpj": "", "monthly_value": None, "start_date": ""},
    )
    assert only_name.status_code == 201, only_name.text
    assert "Nome da pessoa" not in only_name.json()["missing"] and "CNPJ" in only_name.json()["missing"]
    cid = only_name.json()["id"]

    # completo: sem pendência; esvaziar um campo na edição volta a ser pendência (e não é recusado)
    full = client.post("/api/v1/pj", headers=pj_admin, json=contract(7, cost_center_id=st["ccs"]["9001"]))
    assert full.status_code == 201 and full.json()["missing"] == []
    full_id = full.json()["id"]
    upd = client.patch(f"/api/v1/pj/{full_id}", headers=pj_admin, json={"cnpj": "", "monthly_value": ""})
    assert upd.status_code == 200, upd.text
    assert upd.json()["missing"] == ["CNPJ", "Valor mensal"] and upd.json()["cnpj"] is None
    assert client.patch(f"/api/v1/pj/{cid}", headers=pj_admin, json={"cnpj": "11.222.333/0001-82"}).status_code == 422

    # totais: valor mensal vazio soma zero; contadores de pendência e filtro "Com pendência"
    other = client.post("/api/v1/pj", headers=pj_admin, json=contract(8, cost_center_id=st["ccs"]["9001"]))
    assert other.json()["missing"] == []
    s = client.get("/api/v1/pj/summary", headers=pj_admin, params={"year": 2026}).json()
    assert s["active"] == 4 and s["monthly_total"] == "10000.00"  # só o contrato 8 tem valor
    assert s["monthly_missing"] == 3 and s["pending"] == 3
    pending = client.get("/api/v1/pj", headers=pj_admin, params={"pending": "true"}).json()["items"]
    assert len(pending) == 3 and other.json()["id"] not in {i["id"] for i in pending}
    # busca por CNPJ com contratos sem CNPJ
    found = client.get("/api/v1/pj", headers=pj_admin, params={"q": fake_cnpj(8)[:6]}).json()["items"]
    assert [i["id"] for i in found] == [other.json()["id"]]

    # Excel com vazios e a coluna "Falta preencher"
    from openpyxl import load_workbook

    xlsx = client.get("/api/v1/pj/export.xlsx", headers=pj_admin, params={"year": 2026})
    rows = list(load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert rows[0][-1] == "Falta preencher" and len(rows) == 5
    no_name = next(r for r in rows[1:] if r[0] is None)
    assert no_name[12] is None and no_name[9] is None and "CNPJ" in no_name[-1]

    # auditoria continua sem valores em R$
    logs = client.get("/api/v1/audit-logs", headers=pj_admin, params={"entity_type": "pj_contract"}).json()["items"]
    assert "10000" not in json.dumps(logs)


def test_crud_archive_and_audit_without_money(client, pj_admin):
    st = make_structure(client, pj_admin)
    resp = client.post(
        "/api/v1/pj",
        headers=pj_admin,
        json=contract(
            1,
            cnpj=fake_cnpj(1).replace(".", "").replace("/", "").replace("-", ""),
            cost_center_id=st["ccs"]["9001"],
            email="Pessoa.PJ1@Example.com",
            start_date="2026-03-16",
            monthly_value="15432.10",
            annual_bonus="9876.54",
        ),
    )
    assert resp.status_code == 201, resp.text
    item = resp.json()
    assert item["cnpj_formatted"] == fake_cnpj(1) and len(item["cnpj"]) == 14
    assert item["cost_center"] == "9001 · CC PJ 9001" and item["department"] == "Área PJ A"
    assert item["email"] == "pessoa.pj1@example.com"
    assert item["status"] == "ACTIVE" and item["monthly_value"] == "15432.10"
    assert item["tenure"] == {"years": 0, "months": 6}
    assert item["bonus"] == {"year": 2026, "months": 10, "due": "8230.45"}  # 9876,54 × 10 ÷ 12 (março conta)
    cid = item["id"]
    assert client.get(f"/api/v1/pj/{cid}", headers=pj_admin, params={"year": 2027}).json()["bonus"]["months"] == 12

    upd = client.patch(
        f"/api/v1/pj/{cid}",
        headers=pj_admin,
        json={"monthly_value": "16000.00", "role": "Especialista", "end_date": "2026-09-30"},
    )
    assert upd.status_code == 200, upd.text
    assert upd.json()["status"] == "ENDED" and upd.json()["bonus"]["months"] == 7  # março a setembro
    assert upd.json()["tenure"] == {"years": 0, "months": 6}  # até o término
    bad = client.patch(f"/api/v1/pj/{cid}", headers=pj_admin, json={"end_date": "2026-01-01"})
    assert bad.status_code == 422
    assert client.get(f"/api/v1/pj/{cid}", headers=pj_admin).json()["end_date"] == "2026-09-30"  # nada gravado

    assert client.delete(f"/api/v1/pj/{cid}", headers=pj_admin).json() == {"id": cid}
    assert client.get(f"/api/v1/pj/{cid}", headers=pj_admin).status_code == 404
    assert client.patch(f"/api/v1/pj/{cid}", headers=pj_admin, json={"role": "x"}).status_code == 404
    assert client.get("/api/v1/pj", headers=pj_admin).json()["items"] == []

    logs = client.get("/api/v1/audit-logs", headers=pj_admin, params={"entity_type": "pj_contract"}).json()["items"]
    actions = [log["action"] for log in logs]
    assert actions.count("CREATE") == 1 and actions.count("UPDATE") == 1 and actions.count("ARCHIVE") == 1
    assert actions.count("PJ_VIEW") == 1  # consulta registrada no máximo uma vez a cada 30 min
    created = next(log for log in logs if log["action"] == "CREATE")
    assert created["after"]["monthly_value"] == "(sigiloso)" and created["after"]["annual_bonus"] == "(sigiloso)"
    update = next(log for log in logs if log["action"] == "UPDATE")
    assert update["before"]["monthly_value"] == "(sigiloso)" and update["after"]["monthly_value"] == "(alterado)"
    assert update["after"]["role"] == "Especialista" and "annual_bonus" not in update["after"]
    text_ = json.dumps(logs)
    for value in ("15432", "16000", "9876", "8230"):
        assert value not in text_  # nenhum valor em R$ na auditoria


def test_audit_of_pj_is_hidden_without_the_flag(client, pj_admin):
    assert client.post("/api/v1/pj", headers=pj_admin, json=contract()).status_code == 201
    create_user(client, pj_admin, "ctrl@test.com", ["CONTROLLER"])
    controller = login(client, "ctrl@test.com", "Senha@123")
    logs = client.get("/api/v1/audit-logs", headers=controller, params={"entity_type": "pj_contract"}).json()
    assert logs["total"] == 0
    everything = client.get("/api/v1/audit-logs", headers=controller, params={"limit": 1000}).json()["items"]
    assert all(log["entity_type"] != "pj_contract" for log in everything)
    assert (
        client.get("/api/v1/audit-logs", headers=pj_admin, params={"entity_type": "pj_contract"}).json()["total"] >= 1
    )


def test_summary_filters_and_export(client, pj_admin):
    st = make_structure(client, pj_admin)
    cc_a, cc_b = st["ccs"]["9001"], st["ccs"]["9003"]
    bodies = [
        contract(1, cost_center_id=cc_a),  # desde 2020: 12 meses
        contract(2, cost_center_id=cc_a, start_date="2026-03-20", monthly_value="5000", annual_bonus="6000"),  # 10
        contract(3, cost_center_id=cc_b, start_date="2024-01-10", end_date="2026-06-30", monthly_value="8000"),  # 6
        contract(4, cost_center_id=cc_b, annual_bonus=None, monthly_value="7000"),  # sem bonificação
    ]
    for body in bodies:
        assert client.post("/api/v1/pj", headers=pj_admin, json=body).status_code == 201

    s = client.get("/api/v1/pj/summary", headers=pj_admin, params={"year": 2026}).json()
    assert s == {
        "year": 2026,
        "active": 3,
        "ended": 1,
        "monthly_total": "22000.00",  # 10000 + 5000 + 7000 (ativos)
        "annual_bonus_total": "18000.00",  # 12000 + 6000 (ativos)
        "bonus_due_total": "23000.00",  # 12000 + 5000 + 6000 (encerrado no ano conta até junho)
        "monthly_missing": 0,
        "pending": 0,  # todos completos (nome, empresa, CNPJ, função, CC, valor e admissão)
    }
    s2025 = client.get("/api/v1/pj/summary", headers=pj_admin, params={"year": 2025}).json()
    assert s2025["bonus_due_total"] == "24000.00"  # pessoas 1 e 3 o ano todo; 2 ainda não tinha entrado
    assert client.get("/api/v1/pj/summary", headers=pj_admin).json()["year"] == 2026
    sb = client.get("/api/v1/pj/summary", headers=pj_admin, params={"cost_center_id": cc_b}).json()
    assert sb["active"] == 1 and sb["ended"] == 1 and sb["monthly_total"] == "7000.00"
    assert client.get("/api/v1/pj/summary", headers=pj_admin, params={"year": 1500}).status_code == 422

    def names(**params):
        return [i["name"] for i in client.get("/api/v1/pj", headers=pj_admin, params=params).json()["items"]]

    assert names() == ["PESSOA PJ 1", "PESSOA PJ 2", "PESSOA PJ 4", "PESSOA PJ 3"]  # ativos primeiro
    assert names(status="ENDED") == ["PESSOA PJ 3"]
    assert names(status="ACTIVE", cost_center_id=cc_b) == ["PESSOA PJ 4"]
    assert names(q="ficticia 2") == ["PESSOA PJ 2"]
    assert names(q="33.333.3") == ["PESSOA PJ 3"]

    opts = client.get("/api/v1/pj/options", headers=pj_admin).json()
    assert opts["all"] is True
    by_code = {c["code"]: c for c in opts["cost_centers"]}
    assert by_code["9001"]["department"] == "Área PJ A" and by_code["9003"]["department"] == "Área PJ B"

    resp = client.get("/api/v1/pj/export.xlsx", headers=pj_admin, params={"year": 2026})
    assert resp.status_code == 200 and resp.content[:2] == b"PK"
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    assert rows[0][5] == "Área" and rows[0][15] == "Bonificação devida em 2026"
    assert [r[0] for r in rows[1:]] == ["PESSOA PJ 1", "PESSOA PJ 2", "PESSOA PJ 4", "PESSOA PJ 3"]


def test_photos(client, pj_admin):
    cid = client.post("/api/v1/pj", headers=pj_admin, json=contract()).json()["id"]
    put = client.put(f"/api/v1/pj/{cid}/photo", headers=pj_admin, files={"file": ("f.jpg", JPEG, "image/jpeg")})
    assert put.status_code == 200
    photos = client.get("/api/v1/pj/photos", headers=pj_admin).json()
    assert photos[str(cid)].startswith("data:image/jpeg;base64,")
    assert client.get(f"/api/v1/pj/{cid}", headers=pj_admin).json()["has_photo"] is True
    gif = client.put(f"/api/v1/pj/{cid}/photo", headers=pj_admin, files={"file": ("f.gif", b"GIF89a", "image/gif")})
    assert gif.status_code == 422
    big = client.put(
        f"/api/v1/pj/{cid}/photo", headers=pj_admin, files={"file": ("f.jpg", b"0" * (400 * 1024 + 1), "image/jpeg")}
    )
    assert big.status_code == 422
    assert client.patch(f"/api/v1/pj/{cid}", headers=pj_admin, json={"photo_blurred": True}).json()["photo_blurred"]
    assert client.delete(f"/api/v1/pj/{cid}/photo", headers=pj_admin).status_code == 200
    assert client.get("/api/v1/pj/photos", headers=pj_admin).json() == {}
    # contrato arquivado: a foto sai da lista
    client.put(f"/api/v1/pj/{cid}/photo", headers=pj_admin, files={"file": ("f.jpg", JPEG, "image/jpeg")})
    client.delete(f"/api/v1/pj/{cid}", headers=pj_admin)
    assert client.get("/api/v1/pj/photos", headers=pj_admin).json() == {}


# ---------------------------------------------------------------- escopo "Da área"


def test_area_scope(client, pj_admin):
    from openpyxl import load_workbook

    st = make_structure(client, pj_admin)
    ccs = st["ccs"]
    ids = {}
    for n, cc in ((1, ccs["9001"]), (2, ccs["9002"]), (3, ccs["9003"]), (4, None)):
        resp = client.post("/api/v1/pj", headers=pj_admin, json=contract(n, cost_center_id=cc))
        assert resp.status_code == 201
        ids[n] = resp.json()["id"]
    for cid in ids.values():
        client.put(f"/api/v1/pj/{cid}/photo", headers=pj_admin, files={"file": ("f.jpg", JPEG, "image/jpeg")})

    # gestor do CC 9001: a área dele é "Área PJ A" → vê os contratos dos CCs 9001 e 9002, não os de 9003 nem sem CC
    manager_id = create_user(client, pj_admin, "gestor@test.com", ["MANAGER"], manager_of=[ccs["9001"]])
    grant(client, pj_admin, manager_id)  # Da área
    manager = login(client, "gestor@test.com", "Senha@123")

    assert [i["id"] for i in client.get("/api/v1/pj", headers=manager).json()["items"]] == [ids[1], ids[2]]
    assert client.get("/api/v1/pj/summary", headers=manager).json()["active"] == 2
    opts = client.get("/api/v1/pj/options", headers=manager).json()
    assert opts["all"] is False and [c["code"] for c in opts["cost_centers"]] == ["9001", "9002"]
    assert set(client.get("/api/v1/pj/photos", headers=manager).json()) == {str(ids[1]), str(ids[2])}
    for other in (ids[3], ids[4]):
        assert client.get(f"/api/v1/pj/{other}", headers=manager).status_code == 404
        assert client.patch(f"/api/v1/pj/{other}", headers=manager, json={"role": "x"}).status_code == 404
        assert client.delete(f"/api/v1/pj/{other}", headers=manager).status_code == 404
        photo = client.put(f"/api/v1/pj/{other}/photo", headers=manager, files={"file": ("f.jpg", JPEG, "image/jpeg")})
        assert photo.status_code == 404
        assert client.delete(f"/api/v1/pj/{other}/photo", headers=manager).status_code == 404
    xlsx = client.get("/api/v1/pj/export.xlsx", headers=manager)
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert [row[0] for row in ws.iter_rows(min_row=2, values_only=True)] == ["PESSOA PJ 1", "PESSOA PJ 2"]

    # escrita: CC obrigatório e dentro da área
    no_cc = client.post("/api/v1/pj", headers=manager, json=contract(5))
    assert no_cc.status_code == 422 and "Informe o centro de custo" in no_cc.json()["detail"]
    out = client.post("/api/v1/pj", headers=manager, json=contract(5, cost_center_id=ccs["9003"]))
    assert out.status_code == 403 and "fora da sua área" in out.json()["detail"]
    created = client.post("/api/v1/pj", headers=manager, json=contract(5, cost_center_id=ccs["9002"]))
    assert created.status_code == 201
    moved = client.patch(f"/api/v1/pj/{ids[1]}", headers=manager, json={"cost_center_id": ccs["9003"]})
    assert moved.status_code == 403
    cleared = client.patch(f"/api/v1/pj/{ids[1]}", headers=manager, json={"cost_center_id": None})
    assert cleared.status_code == 422
    assert client.get(f"/api/v1/pj/{ids[1]}", headers=manager).json()["cost_center_id"] == ccs["9001"]  # nada gravado
    assert client.patch(f"/api/v1/pj/{ids[1]}", headers=manager, json={"role": "Revisor"}).status_code == 200

    # Controladoria com "Da área": vale só a própria área, mesmo que o perfil veja todos os CCs do Orçamento
    ctrl_id = create_user(client, pj_admin, "ctrl@test.com", ["CONTROLLER"])
    grant(client, pj_admin, ctrl_id)
    controller = login(client, "ctrl@test.com", "Senha@123")
    assert client.get("/api/v1/pj", headers=controller).json()["items"] == []  # sem CC próprio: nada
    assert client.get("/api/v1/pj/options", headers=controller).json()["cost_centers"] == []
    assert client.get("/api/v1/audit-logs", headers=controller, params={"entity_type": "pj_contract"}).json() == {
        "total": 0,
        "items": [],
    }
    # com um escopo de CC na área B, passa a ver só a área B (inclusive a auditoria desses contratos)
    client.put(f"{USERS}/{ctrl_id}/scopes", headers=pj_admin, json=[{"cost_center_id": ccs["9003"]}])
    assert [i["id"] for i in client.get("/api/v1/pj", headers=controller).json()["items"]] == [ids[3]]
    seen = client.get("/api/v1/audit-logs", headers=controller, params={"entity_type": "pj_contract"}).json()["items"]
    assert seen and {log["entity_id"] for log in seen if log["action"] != "PJ_VIEW"} == {str(ids[3])}

    # o escopo do gestor, conferido pelo serviço (ele não lê a auditoria)
    mine = {str(ids[1]), str(ids[2]), str(created.json()["id"])}
    with SessionLocal() as db:
        gestor = db.get(User, manager_id)
        assert area_cost_center_ids(db, gestor) == {ccs["9001"], ccs["9002"]}
        assert set(audit_contract_ids(db, gestor)) == mine
        assert audit_contract_ids(db, db.query(User).filter_by(email="admin@test.com").one()) is None

    # "Todos" vê tudo, inclusive sem CC
    grant(client, pj_admin, manager_id, all_=True)
    assert len(client.get("/api/v1/pj", headers=manager).json()["items"]) == 5
    assert client.get(f"/api/v1/pj/{ids[4]}", headers=manager).status_code == 200


# ---------------------------------------------------------------- migração


def test_migrations_0009_0011_up_down_without_drift():
    from alembic.autogenerate import compare_metadata
    from alembic.config import Config
    from alembic.migration import MigrationContext

    from alembic import command

    backend = Path(__file__).resolve().parent.parent
    cfg = Config(str(backend / "alembic.ini"))
    cfg.set_main_option("script_location", str(backend / "alembic"))
    Base.metadata.drop_all(engine)
    with engine.begin() as conn:
        conn.execute(text("drop table if exists alembic_version"))
    try:
        command.upgrade(cfg, "head")
        with engine.connect() as conn:
            diff = compare_metadata(MigrationContext.configure(conn, opts={"compare_type": True}), Base.metadata)
        pj_diff = [d for d in diff if "pj_" in repr(d) or "can_view" in repr(d)]
        assert pj_diff == [], pj_diff
        assert diff == [], diff
        # 0011 para baixo: vazios recebem '' / 0 / data de criação antes de voltar a NOT NULL
        with engine.begin() as conn:
            conn.execute(text("insert into pj_contracts default values"))
        command.downgrade(cfg, "0010")
        cols = {c["name"]: c["nullable"] for c in inspect(engine).get_columns("pj_contracts")}
        assert not any(cols[c] for c in ("name", "company_name", "cnpj", "monthly_value", "start_date"))
        with engine.connect() as conn:
            row = conn.execute(text("select name, cnpj, monthly_value, start_date from pj_contracts")).one()
        assert row[0] == "" and row[1] == "" and row[2] == 0 and row[3] is not None
        command.upgrade(cfg, "0011")
        cols = {c["name"]: c["nullable"] for c in inspect(engine).get_columns("pj_contracts")}
        assert all(cols[c] for c in ("name", "company_name", "cnpj", "monthly_value", "start_date"))
        command.downgrade(cfg, "0008")
        insp = inspect(engine)
        assert "pj_contracts" not in insp.get_table_names() and "pj_photos" not in insp.get_table_names()
        assert "can_view_pj" not in {c["name"] for c in insp.get_columns("users")}
        command.upgrade(cfg, "head")
        assert "pj_photos" in inspect(engine).get_table_names()
    finally:
        with engine.begin() as conn:
            conn.execute(text("drop table if exists alembic_version"))


# ---------------------------------------------------------------- orçamento


def test_contracts_feed_the_budget_as_a_single_account(client, admin, pj_admin):
    """Contratos PJ entram no orçamento do ciclo (09/10/2026): mensal × meses de vigência com IPCA (pj.adjustment_pct,
    3,5% no seed) + bonificação anual cheia diluída nos meses (paga no ano seguinte), na conta pj.budget_account
    (6010201016), só o total por CC × mês. Sem CC, encerrado antes do ano ou arquivado: fora."""
    st = make_structure(client, admin)
    cycle_id = client.get("/api/v1/cycles", headers=admin).json()[0]["id"]
    client.post(f"/api/v1/cycles/{cycle_id}/open", headers=admin)

    def opex():
        ov = client.get("/api/v1/consolidation/overview", headers=admin).json()
        return Decimal(ov["modules"]["OPEX"]["proposed"])

    before = opex()
    cc = st["ccs"]["9001"]
    assert client.post("/api/v1/pj", headers=pj_admin, json=contract(1, cost_center_id=cc)).status_code == 201
    mid = contract(2, cost_center_id=cc, monthly_value="1000.00", annual_bonus="1200.00", start_date="2027-07-10")
    assert client.post("/api/v1/pj", headers=pj_admin, json=mid).status_code == 201
    assert client.post("/api/v1/pj", headers=pj_admin, json=contract(3)).status_code == 201  # sem CC
    ended = contract(4, cost_center_id=cc, start_date="2020-01-01", end_date="2026-12-31")
    assert client.post("/api/v1/pj", headers=pj_admin, json=ended).status_code == 201
    # 10.000 × 1,035 × 12 + 12.000 (cheia) = 136.200; 1.000 × 1,035 × 6 (jul–dez) + 1.200 (cheia) = 7.410
    assert opex() - before == Decimal("143610.00")
