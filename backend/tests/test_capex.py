import io

from openpyxl import load_workbook

from tests import builders
from tests.test_imports import import_and_load, status, upload
from tests.test_opex import CC, _user


def _cc(client, admin):
    return next(c for c in client.get("/api/v1/cost-centers", headers=admin).json() if c["code"] == CC)


def test_capex_manual_flow(client, admin, run_worker):
    import_and_load(client, admin, run_worker, builders.opex_template_bd(), "bd.xlsx", dataset_type="MASTER_DATA")
    mgr_id, mgr = _user(client, admin, "gestor@t.com", ["MANAGER"])
    cc = _cc(client, admin)
    client.patch(f"/api/v1/cost-centers/{cc['id']}", headers=admin, json={"manager_user_id": mgr_id})
    cycle_id = client.get("/api/v1/cycles", headers=admin).json()[0]["id"]
    client.post(f"/api/v1/cycles/{cycle_id}/open", headers=admin)

    opts = client.get("/api/v1/capex/options", headers=mgr).json()
    acc = {a["code"]: a["id"] for a in opts["accounts"]}
    assert "1020601005" in acc and "6010301002" not in acc  # só contas de ativo
    assert len(opts["lookups"]["CAPEX_PROJECT_TYPE"]) == 8

    head = client.get(f"/api/v1/capex/cost-centers/{cc['id']}", headers=mgr).json()
    sub = head["submission_id"]
    assert head["status"] == "DRAFT" and head["permissions"]["edit"]

    # projeto sem tipo → pendência crítica; cronograma divergente → crítica
    project = client.post(
        f"/api/v1/capex/submissions/{sub}/projects",
        headers=mgr,
        json={
            "title": "Modernização do BI",
            "is_project": True,
            "items": [
                {
                    "account_id": acc["1020601005"],
                    "item_name": "Servidor",
                    "unit_value": 25000,
                    "quantity": 2,
                    "useful_life_months": 60,
                    "values": {"4": 30000},
                }
            ],
        },
    )
    assert project.status_code == 201, project.text
    p = project.json()
    assert p["code"] == "CPX-001" and p["total"] == "50000.00"
    codes = {i["code"] for i in p["issues"]} | {i["code"] for i in p["items"][0]["issues"]}
    assert {"CAPEX_NO_PROJECT_TYPE", "CAPEX_NO_JUSTIFICATION", "CAPEX_SCHEDULE_MISMATCH"} <= codes

    resp = client.post(f"/api/v1/capex/submissions/{sub}/actions/submit", headers=mgr, json={})
    assert resp.status_code == 409 and "difere do valor total" in resp.json()["detail"]

    # conta de despesa é recusada
    opex_acc = client.get("/api/v1/opex/options", headers=mgr).json()["accounts"][0]["id"]
    bad = client.post(
        f"/api/v1/capex/projects/{p['id']}/items",
        headers=mgr,
        json={"account_id": opex_acc, "item_name": "X", "unit_value": 5000, "quantity": 1},
    )
    assert bad.status_code == 422

    item_id = p["items"][0]["id"]
    fixed = client.patch(f"/api/v1/capex/items/{item_id}", headers=mgr, json={"values": {"5": 20000}}).json()
    assert fixed["scheduled"] == "50000.00" and fixed["difference"] == "0.00"
    client.patch(
        f"/api/v1/capex/projects/{p['id']}",
        headers=mgr,
        json={"project_type_code": "Automação e Transformação Digital", "justification": "Reduz 2 FTE"},
    )
    # aquisição avulsa de baixo valor: só aviso, não bloqueia
    small = client.post(
        f"/api/v1/capex/submissions/{sub}/projects",
        headers=mgr,
        json={
            "title": "Cadeiras",
            "justification": "Reposição",
            "items": [
                {
                    "account_id": acc["1020601004"],
                    "item_name": "Cadeira",
                    "unit_value": 900,
                    "quantity": 5,
                    "values": {"2": 4500},
                }
            ],
        },
    ).json()
    assert [i["code"] for i in small["items"][0]["issues"]] == ["CAPEX_BELOW_MIN_VALUE"]

    view = client.get(f"/api/v1/capex/submissions/{sub}/view", headers=mgr).json()
    assert view["totals"]["proposed"] == "54500.00" and view["issues"]["critical"] == 0
    assert view["monthly"][3] == "30000.00" and view["monthly"][4] == "20000.00"
    assert {r["label"] for r in view["by_type"]} == {"Automação e Transformação Digital", "Aquisição avulsa"}

    head = client.post(f"/api/v1/capex/submissions/{sub}/actions/submit", headers=mgr, json={}).json()
    assert head["status"] == "SUBMITTED" and not head["permissions"]["edit"]
    assert client.patch(f"/api/v1/capex/items/{item_id}", headers=mgr, json={"quantity": 3}).status_code == 409
    client.post(f"/api/v1/capex/submissions/{sub}/actions/start_review", headers=admin, json={})
    head = client.post(f"/api/v1/capex/submissions/{sub}/actions/approve", headers=admin, json={}).json()
    assert head["status"] == "APPROVED"

    summary = client.get("/api/v1/capex/summary", headers=admin).json()
    row = next(r for r in summary["rows"] if r["code"] == CC)
    assert row["total"] == "54500.00" and row["requests"] == 2 and row["status"] == "APPROVED"
    assert summary["by_account"][0]["code"] == "1020601005"

    # outro gestor não acessa; OPEX e CAPEX do mesmo CC são orçamentos separados
    _, other = _user(client, admin, "outro@t.com", ["MANAGER"])
    assert client.get(f"/api/v1/capex/submissions/{sub}/view", headers=other).status_code == 403
    opex = client.get(f"/api/v1/opex/cost-centers/{cc['id']}", headers=mgr).json()
    assert opex["submission_id"] != sub and opex["status"] == "DRAFT"


def test_capex_template_import(client, admin, run_worker):
    batch_id = upload(client, admin, builders.capex_template_filled(), "Template_CAPEX 2027_Gestor.xlsx")
    run_worker()
    b = status(client, admin, batch_id)
    assert b["dataset_type"] == "CAPEX_TEMPLATE", b
    assert b["summary"]["meta"]["parts"] == {"master": 5, "asset_items": 3, "capex_items": 4}
    cmp = b["summary"]["comparison"]
    assert cmp["catalog"] == {"new": 3, "existing": 0}
    assert cmp["budget"][0]["lines"] == 3 and cmp["budget"][0]["total"] == "34000.00"
    errors = client.get(f"/api/v1/imports/{batch_id}/errors", headers=admin).json()
    codes = {e["code"] for e in (errors["items"] if isinstance(errors, dict) else errors)}
    assert {"WRONG_NATURE", "CAPEX_SCHEDULE_MISMATCH", "CAPEX_BELOW_MIN_VALUE", "NEW_ASSET_CLASS"} <= codes
    assert "CAPEX_SOFTWARE" in codes  # licença de BI na conta de software: avaliar se é assinatura (OPEX)

    assert client.post(f"/api/v1/imports/{batch_id}/confirm", headers=admin).status_code == 200
    run_worker()
    final = status(client, admin, batch_id)
    assert final["status"] == "COMPLETED", final
    load = final["summary"]["load"]
    assert load["capex"] == {"requests_created": 2, "items_created": 3, "cost_centers": 1}
    assert load["catalog"]["items_created"] == 3 and load["catalog"]["classes_created"] == 1

    cc = _cc(client, admin)
    head = client.get(f"/api/v1/capex/cost-centers/{cc['id']}", headers=admin).json()
    assert head["status"] == "IN_PROGRESS"
    view = client.get(f"/api/v1/capex/submissions/{head['submission_id']}/view", headers=admin).json()
    project = next(p for p in view["projects"] if p["is_project"])
    assert project["project_type_code"] == "Automação e Transformação Digital"
    assert [i["item_name"] for i in project["items"]] == ["NOTEBOOK", "LICENÇA"]
    assert project["items"][0]["asset_item_id"] is not None
    assert view["totals"]["proposed"] == "34000.00"
    assert view["issues"]["critical"] == 1  # cadeira: cronograma 3.000 × total 4.000
    # página Apontamentos: as pendências do CC aparecem item a item, críticas primeiro
    found = client.get("/api/v1/findings", headers=admin).json()
    mine = [i for i in found["items"] if i["cost_center_id"] == cc["id"] and i["module"] == "CAPEX"]
    assert mine[0]["severity"] == "CRITICAL" and mine[0]["kind"] == "CAPEX_SCHEDULE_MISMATCH"
    assert mine[0]["subject"].endswith("· CADEIRA") and mine[0]["link"] == f"/capex/{cc['id']}"
    assert {"CAPEX_SOFTWARE", "CAPEX_BELOW_MIN_VALUE"} <= {i["kind"] for i in mine}
    assert found["counts"]["critical"] >= 1 and found["counts"]["cost_centers"] >= 1
    _, outsider = _user(client, admin, "sem-cc@t.com", ["MANAGER"])
    assert client.get("/api/v1/findings", headers=outsider).json()["items"] == []  # gestor só vê os seus CCs

    # CAPEX sem cronograma é orçamento: a cadeira (total 4.000, cronograma 3.000) entra inteira no total do ano,
    # a diferença aparece como "sem cronograma" (Consolidação e Painel), e some quando o cronograma fecha
    ov = client.get("/api/v1/consolidation/overview", headers=admin).json()["modules"]["CAPEX"]
    assert (ov["proposed"], ov["unscheduled"]) == ("34000.00", "1000.00")
    assert sum(float(v) for v in ov["monthly"]) == 33000.0
    panel = client.get("/api/v1/dashboard/overview?years=2027", headers=admin).json()
    assert (panel["kpis"]["ref_ytd"], panel["kpis"]["budget_unscheduled"]) == ("34000.00", "1000.00")
    assert sum(float(r["budget"]) for r in panel["monthly"]) == 33000.0
    capex_card = next(m for m in panel["by_module"] if m["module"] == "CAPEX")
    assert (capex_card["main"], capex_card["unscheduled"]) == ("34000.00", "1000.00")
    assert (
        client.get("/api/v1/dashboard/overview?years=2027&months=5", headers=admin).json()["kpis"]["ref_ytd"]
        == "3000.00"
    )

    # correção na própria página: cronograma da cadeira (4 × R$ 1.000) fecha com o total
    cadeira = mine[0]
    assert cadeira["fix"]["type"] == "schedule" and cadeira["fix"]["total"] == "4000.00" and cadeira["editable"]
    assert client.post("/api/v1/findings/fix", headers=outsider, json={"key": cadeira["key"]}).status_code == 403
    short = client.post("/api/v1/findings/fix", headers=admin, json={"key": cadeira["key"], "values": {"5": 3000}})
    assert short.status_code == 422 and "precisa fechar" in short.json()["detail"]
    fixed = client.post(
        "/api/v1/findings/fix", headers=admin, json={"key": cadeira["key"], "values": {"5": 2000, "6": 2000}}
    )
    assert fixed.status_code == 200, fixed.text
    assert fixed.json()["note"] == "Cronograma: MAI R$ 2.000,00, JUN R$ 2.000,00"
    ov = client.get("/api/v1/consolidation/overview", headers=admin).json()["modules"]["CAPEX"]
    assert (ov["proposed"], ov["unscheduled"]) == ("34000.00", "0.00")
    found = client.get("/api/v1/findings", headers=admin).json()
    assert cadeira["key"] not in {i["key"] for i in found["items"]}
    assert (found["reviews"][0]["action"], found["reviews"][0]["kind"]) == ("CORRECTED", "CAPEX_SCHEDULE_MISMATCH")
    assert found["cost_centers"][0]["capex_submission_id"] == head["submission_id"]

    # aviso mantido pela Controladoria com o motivo; sai dos pendentes e entra no relatório; reabrir devolve
    software = next(i for i in found["items"] if i["kind"] == "CAPEX_SOFTWARE")
    blank = client.post("/api/v1/findings/keep", headers=admin, json={"key": software["key"], "note": " "})
    assert blank.status_code == 422
    kept = client.post(
        "/api/v1/findings/keep", headers=admin, json={"key": software["key"], "note": "Licença perpétua"}
    )
    assert kept.status_code == 200, kept.text
    found = client.get("/api/v1/findings", headers=admin).json()
    assert software["key"] not in {i["key"] for i in found["items"]}
    assert (found["reviews"][0]["action"], found["reviews"][0]["note"]) == ("KEPT", "Licença perpétua")
    report = client.get("/api/v1/findings/export.xlsx", headers=admin)
    assert report.status_code == 200
    assert report.headers["content-disposition"].endswith('filename="Apontamentos_2027_Rev0.xlsx"')
    wb = load_workbook(io.BytesIO(report.content))
    assert wb.sheetnames == ["Resumo", "Pendentes", "Corrigidos e mantidos"]
    done = wb["Corrigidos e mantidos"]
    assert {done.cell(r, 7).value for r in range(2, done.max_row + 1)} == {"Corrigido", "Mantido"}
    assert client.delete(f"/api/v1/findings/reviews/{found['reviews'][0]['id']}", headers=admin).status_code == 200
    assert software["key"] in {i["key"] for i in client.get("/api/v1/findings", headers=admin).json()["items"]}
    licence = project["items"][1]
    assert [i["code"] for i in licence["issues"]] == ["CAPEX_SOFTWARE"]
    # notebook levado para a conta de software: diverge do catálogo (aviso, não bloqueia)
    software_id = next(
        a["id"]
        for a in client.get("/api/v1/capex/options", headers=admin).json()["accounts"]
        if a["code"] == "1020701002"
    )
    moved = client.patch(
        f"/api/v1/capex/items/{project['items'][0]['id']}", headers=admin, json={"account_id": software_id}
    ).json()
    assert [(i["code"], i["severity"]) for i in moved["issues"]] == [("CAPEX_ACCOUNT_MISMATCH", "WARNING")]
    # e volta para a conta do catálogo direto no apontamento
    mismatch = next(
        i
        for i in client.get("/api/v1/findings", headers=admin).json()["items"]
        if i["kind"] == "CAPEX_ACCOUNT_MISMATCH"
    )
    assert mismatch["fix"]["account"] == "1020601005 · Equipamentos de Informática"
    back = client.post("/api/v1/findings/fix", headers=admin, json={"key": mismatch["key"]})
    assert back.status_code == 200 and back.json()["note"].startswith("Conta 1020701002")

    opts = client.get("/api/v1/capex/options", headers=admin).json()
    notebook = next(a for a in opts["asset_items"] if a["name"] == "NOTEBOOK")
    assert notebook["asset_class"] == "Equipamentos de Informática" and notebook["account_id"]

    # solicitação digitada no sistema é preservada; reimportar substitui só as do template
    client.post(
        f"/api/v1/capex/submissions/{head['submission_id']}/projects",
        headers=admin,
        json={"title": "Manual", "justification": "x"},
    )
    again = import_and_load(client, admin, run_worker, builders.capex_template_filled(), "v2.xlsx", force=True)
    assert again["summary"]["load"]["capex"]["requests_replaced"] == 2
    view = client.get(f"/api/v1/capex/submissions/{head['submission_id']}/view", headers=admin).json()
    assert len(view["projects"]) == 3
