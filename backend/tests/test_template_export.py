"""Exportação do orçamento do CC no layout do template e reimportação (ida e volta)."""

from decimal import Decimal

from openpyxl import load_workbook

from app.imports.base import load_sheets
from app.imports.parsers.capex_template import parse_capex_template
from app.imports.parsers.opex_template import parse_opex_template
from tests import builders
from tests.test_imports import import_and_load, status, upload
from tests.test_opex import CC, _user

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _cc(client, admin):
    return next(c for c in client.get("/api/v1/cost-centers", headers=admin).json() if c["code"] == CC)


def test_opex_template_export_roundtrip(client, admin, run_worker):
    import_and_load(client, admin, run_worker, builders.opex_template_filled(), "template.xlsx")
    cc = _cc(client, admin)
    head = client.get(f"/api/v1/opex/cost-centers/{cc['id']}", headers=admin).json()
    sub = head["submission_id"]
    # linha digitada no sistema (sem origem template) também vai para a planilha
    acc = next(
        a for a in client.get("/api/v1/opex/options", headers=admin).json()["accounts"] if a["code"] == "6010301005"
    )
    client.post(
        f"/api/v1/opex/submissions/{sub}/lines",
        headers=admin,
        json={"account_id": acc["id"], "values": {"1": 100}, "description": "Digitada", "supplier": "Loja"},
    )
    import io

    resp = client.get(f"/api/v1/opex/submissions/{sub}/template.xlsx", headers=admin)
    assert resp.status_code == 200 and resp.headers["content-type"] == XLSX
    assert resp.headers["content-disposition"].endswith(f'filename="Template_OPEX_2027_{CC}_Rev0.xlsx"')
    content = resp.content

    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames[:3] == ["Instruções", "BD-Novo", "I - Viagens"] and "VI - DTI" in wb.sheetnames
    viagens = wb["I - Viagens"]
    purposes = {viagens.cell(r, 9).value for r in range(7, 9)}
    assert purposes == {"Auditoria SP", "Visita base Belém"}

    # o importador lê a planilha exportada como template OPEX e encontra os mesmos valores
    sheets = load_sheets(content, "export.xlsx")
    result = parse_opex_template(sheets, {})
    assert result.meta["system_export"] == {"module": "OPEX", "cost_center": CC, "version": "Rev0"}
    assert result.meta["parts"]["budget_lines"] == 5  # 2 viagens + consultoria + DTI + linha digitada
    assert result.meta["consolidator_check"]["I - Viagens"]["ok"] is True
    budget = [r for r in result.records if r.record_type in ("BUDGET_LINE", "TRAVEL")]
    assert not any(r.issues for r in budget), [r.issues for r in budget]
    total = sum((Decimal(v) for r in budget for v in r.data["values"].values()), Decimal(0))
    assert total == Decimal("20500")
    consult = next(r for r in budget if r.data.get("supplier") == "KPMG")
    assert consult.data["account"] == "6010201003" and consult.data["justification"] == "Reajuste IPCA"
    assert consult.data["cost_center"] == CC and consult.data["branch"] == "0001"
    assert consult.data["contract_manager"] == "Ana"  # colunas extras preservadas na ida e volta
    travel = next(r for r in budget if r.record_type == "TRAVEL")
    assert travel.data["company"] == "1001"  # empresa lida da CHAVE, não fixa

    # evento calculado pelo sistema: vai para a aba de leitura e sobrevive à reimportação
    ev_acc = next(
        a for a in client.get("/api/v1/opex/options", headers=admin).json()["accounts"] if a["code"] == "6010501002"
    )
    ev = client.post(
        f"/api/v1/opex/submissions/{sub}/lines",
        headers=admin,
        json={
            "line_type": "EVENT",
            "account_id": ev_acc["id"],
            "description": "Confraternização",
            "event": {"event_type": "Interno", "month": 12, "people": 10, "structure": 1000},
        },
    )
    assert ev.status_code == 201, ev.text
    content = client.get(f"/api/v1/opex/submissions/{sub}/template.xlsx", headers=admin).content
    wb = load_workbook(io.BytesIO(content))
    assert "Eventos (leitura)" in wb.sheetnames and wb["Eventos (leitura)"].cell(7, 4).value == "Confraternização"
    result = parse_opex_template(load_sheets(content, "export.xlsx"), {})
    assert result.meta["parts"]["budget_lines"] == 5  # o evento não entra como linha de template

    # reimportar a exportação substitui TODOS os lançamentos do CC (sem duplicar a linha digitada)
    final = import_and_load(client, admin, run_worker, content, "export.xlsx", force=True)
    assert final["summary"]["load"]["budget"] == {
        "lines_replaced": 5,
        "lines_created": 5,
        "cost_centers": 1,
        "trips": 2,
    }
    lines = client.get(f"/api/v1/opex/submissions/{sub}/lines", headers=admin).json()
    assert len(lines) == 6  # 5 reimportadas + evento preservado
    event_line = next(line for line in lines if line["line_type"] == "EVENT")
    assert event_line["description"] == "Confraternização"
    assert sum(Decimal(line["total"]) for line in lines) == Decimal("20500.00") + Decimal(event_line["total"])
    typed = next(line for line in lines if line["supplier"] == "Loja")
    assert typed["description"] == "Digitada" and typed["values"]["1"] == "100.00"
    consult_line = next(line for line in lines if line["supplier"] == "KPMG")
    assert consult_line["contract_manager"] == "Ana"

    # planilha exportada com uma linha inválida: nada do CC é substituído (senão a linha sumiria)
    wb = load_workbook(io.BytesIO(content))
    dti = wb["VI - DTI"]
    dti.cell(7, 2).value = f"1001-0001-{CC}-9999999999"  # conta inexistente na linha de DTI (CHAVE e código)
    dti.cell(7, 14).value = "9999999999"
    broken = io.BytesIO()
    wb.save(broken)
    batch_id = upload(client, admin, broken.getvalue(), "quebrada.xlsx")
    run_worker()
    b = status(client, admin, batch_id)
    codes = {e["code"] for e in client.get(f"/api/v1/imports/{batch_id}/errors", headers=admin).json()}
    assert "EXPORT_HAS_ERRORS" in codes, (b["status"], codes)
    assert b["summary"]["comparison"]["budget"][0]["replaces_lines"] == 0

    # gestor de outro CC não baixa
    _, other = _user(client, admin, "outro@t.com", ["MANAGER"])
    assert client.get(f"/api/v1/opex/submissions/{sub}/template.xlsx", headers=other).status_code == 403


def test_capex_template_export_roundtrip(client, admin, run_worker):
    import_and_load(client, admin, run_worker, builders.opex_template_bd(), "bd.xlsx", dataset_type="MASTER_DATA")
    final = import_and_load(client, admin, run_worker, builders.capex_template_filled(), "capex.xlsx", force=True)
    cc = _cc(client, admin)
    head = client.get(f"/api/v1/capex/cost-centers/{cc['id']}", headers=admin).json()
    sub = head["submission_id"]
    before = client.get(f"/api/v1/capex/submissions/{sub}/view", headers=admin).json()
    n_items = sum(len(p["items"]) for p in before["projects"])
    assert n_items == final["summary"]["load"]["capex"]["items_created"]

    resp = client.get(f"/api/v1/capex/submissions/{sub}/template.xlsx", headers=admin)
    assert resp.status_code == 200
    assert resp.headers["content-disposition"].endswith(f'filename="Template_CAPEX_2027_{CC}_Rev0.xlsx"')
    content = resp.content
    result = parse_capex_template(load_sheets(content, "export.xlsx"), {})
    assert result.meta["system_export"]["cost_center"] == CC and result.meta["parts"]["capex_items"] == n_items
    items = [r for r in result.records if r.record_type == "CAPEX_ITEM"]
    assert not any(r.issues for r in items), [r.issues for r in items]
    assert {r.data["item"] for r in items} == {"NOTEBOOK", "LICENÇA", "CADEIRA"}
    notebook = next(r for r in items if r.data["item"] == "NOTEBOOK")
    assert notebook.data["values"] == {3: "18000"} and notebook.data["is_project"] is True
    assert notebook.data["project_type"] == "Automação e Transformação Digital"

    final2 = import_and_load(client, admin, run_worker, content, "export.xlsx", force=True)
    load = final2["summary"]["load"]["capex"]
    assert load["items_created"] == n_items and load["requests_replaced"] == len(before["projects"])
    after = client.get(f"/api/v1/capex/submissions/{sub}/view", headers=admin).json()
    assert len(after["projects"]) == len(before["projects"])
    assert {p["title"] for p in after["projects"]} == {p["title"] for p in before["projects"]}
    assert result.meta["system_export"]["module"] == "CAPEX"
    assert after["totals"]["proposed"] == before["totals"]["proposed"] and Decimal(after["totals"]["proposed"]) > 0
