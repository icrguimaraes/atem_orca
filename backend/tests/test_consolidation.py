import io
from decimal import Decimal

from openpyxl import load_workbook

from tests import builders
from tests.test_imports import import_and_load
from tests.test_opex import CC, HIST_2026
from tests.test_personnel import QUADRO


def _setup(client, admin, run_worker):
    import_and_load(client, admin, run_worker, builders.opex_template_bd(), "bd.xlsx", dataset_type="MASTER_DATA")
    import_and_load(client, admin, run_worker, builders.realizado_wide(HIST_2026), "r26.xlsx")
    import_and_load(client, admin, run_worker, builders.quadro_funcionarios(QUADRO), "quadro.xlsx")
    cycle_id = client.get("/api/v1/cycles", headers=admin).json()[0]["id"]
    client.post(f"/api/v1/cycles/{cycle_id}/open", headers=admin)
    cc = next(c for c in client.get("/api/v1/cost-centers", headers=admin).json() if c["code"] == CC)
    opex = client.get(f"/api/v1/opex/cost-centers/{cc['id']}", headers=admin).json()["submission_id"]
    acc = {a["code"]: a["id"] for a in client.get("/api/v1/opex/options", headers=admin).json()["accounts"]}
    line = client.post(
        f"/api/v1/opex/submissions/{opex}/lines",
        headers=admin,
        json={"account_id": acc["6010301002"], "values": {m: 450 for m in range(1, 13)}},
    ).json()[0]
    capex = client.get(f"/api/v1/capex/cost-centers/{cc['id']}", headers=admin).json()["submission_id"]
    capex_acc = {a["code"]: a["id"] for a in client.get("/api/v1/capex/options", headers=admin).json()["accounts"]}
    client.post(
        f"/api/v1/capex/submissions/{capex}/projects",
        headers=admin,
        json={
            "title": "Servidores",
            "justification": "Capacidade",
            "items": [
                {
                    "account_id": capex_acc["1020601005"],
                    "item_name": "Servidor",
                    "unit_value": 25000,
                    "quantity": 2,
                    "values": {"4": 50000},
                }
            ],
        },
    )
    return cc, opex, line


def test_overview_export_freeze_and_revision(client, admin, run_worker):
    cc, opex_sub, line = _setup(client, admin, run_worker)
    ov = client.get("/api/v1/consolidation/overview", headers=admin).json()
    m = ov["modules"]
    assert (m["OPEX"]["proposed"], m["CAPEX"]["proposed"], m["PERSONNEL"]["proposed"]) == (
        "5400.00",
        "50000.00",
        "469350.00",
    )
    assert ov["total"] == "524750.00" and ov["version"]["label"] == "1.0"
    assert m["OPEX"]["ref_annualized"] == "18000.00"  # (1000 + 500) × 8 meses → anualizado
    row = next(r for r in ov["matrix"] if r["code"] == CC)
    assert row["status"] == {"OPEX": "IN_PROGRESS", "CAPEX": "IN_PROGRESS", "PERSONNEL": "IN_PROGRESS"}
    var = {v["account"]: v for v in ov["variations"]}
    assert var["6010301001"]["flags"] == ["NO_BUDGET"] and var["6010301002"]["variation"] == "-600.00"
    assert var["6010101001"]["proposed"] == "267750.00"  # salário (com reajuste) do quadro
    # parte do multiplicador: rateada entre as contas do parâmetro padrão `personnel.charges_split`
    from app.seed_data import CYCLE_PARAMETERS

    split_codes = set(CYCLE_PARAMETERS["personnel.charges_split"][0])
    charges = [v for v in var.values() if v["account"] in split_codes and v["proposed"] != "0.00"]
    assert {v["account"] for v in charges} == split_codes
    assert sum(Decimal(v["proposed"]) for v in charges) == Decimal("201600.00")
    pessoas = next(p for p in ov["by_package"] if p["label"] == "Pessoas")
    assert pessoas["proposed"] == "469350.00"

    points = client.get("/api/v1/consolidation/attention-points", headers=admin).json()
    kinds = {p["kind"] for p in points["points"]}
    assert {"ACCOUNT_CONFIG", "JUSTIFICATION"} <= kinds  # contas de pessoal não cadastradas; vaga sem justificativa

    resp = client.get("/api/v1/consolidation/export.xlsx", headers=admin)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == [
        "Resumo",
        "Carga SAP",
        "Consolidado",
        "Variações por conta",
        "OPEX (linhas)",
        "CAPEX (itens)",
        "Pessoal (quadro)",
        "Status por CC",
    ]
    sap = wb["Carga SAP"]
    keys = [sap.cell(r, 5).value for r in range(2, sap.max_row) if sap.cell(r, 5).value]
    assert f"1001--{CC}-6010101001" in keys and f"1001--{CC}-1020601005" in keys
    # coluna 19 = Total (a 18 é "Sem cronograma")
    totals = [sap.cell(r, 19).value for r in range(2, sap.max_row) if isinstance(sap.cell(r, 19).value, (int, float))]
    assert round(sum(totals), 2) == 524750.00
    assert wb["Pessoal (quadro)"].max_row >= 5 and wb["OPEX (linhas)"].cell(2, 4).value == "6010301002"

    # gestor não congela; Controladoria congela e a versão vira só leitura
    from tests.test_opex import _user

    _, mgr = _user(client, admin, "gestor@t.com", ["MANAGER"])
    assert client.post("/api/v1/consolidation/freeze", headers=mgr, json={}).status_code == 403
    fz = client.post("/api/v1/consolidation/freeze", headers=admin, json={"reason": "Fechamento"}).json()
    assert fz["version"] == "1.0" and fz["total"] == "524750.00"
    resp = client.patch(f"/api/v1/opex/lines/{line['id']}", headers=admin, json={"values": {"1": 999}})
    assert resp.status_code == 409 and "congelada" in resp.json()["detail"]
    resp = client.post(f"/api/v1/opex/submissions/{opex_sub}/actions/submit", headers=admin, json={})
    assert resp.status_code == 409
    assert client.post("/api/v1/consolidation/freeze", headers=admin, json={}).status_code == 409

    # o quadro muda depois do congelamento: a versão congelada não muda (fotografia)
    import_and_load(
        client,
        admin,
        run_worker,
        builders.quadro_funcionarios(QUADRO[:1]),
        "quadro2.xlsx",
        force=True,
        deactivate_missing="true",
    )
    assert client.get("/api/v1/consolidation/overview", headers=admin).json()["total"] == "524750.00"

    # revisão: 1.1 em elaboração com tudo copiado; 1.0 continua consultável
    assert client.post("/api/v1/consolidation/revise", headers=admin, json={"reason": " "}).status_code == 409
    rv = client.post("/api/v1/consolidation/revise", headers=admin, json={"reason": "Ajuste de pessoal"}).json()
    assert rv["version"] == "1.1" and [v["status"] for v in rv["versions"]] == ["FROZEN", "WORKING"]
    head = client.get(f"/api/v1/opex/cost-centers/{cc['id']}", headers=admin).json()
    assert head["version"] == "1.1" and head["submission_id"] != opex_sub and head["permissions"]["edit"]
    lines = client.get(f"/api/v1/opex/submissions/{head['submission_id']}/lines", headers=admin).json()
    assert len(lines) == 1 and lines[0]["total"] == "5400.00"
    capex_head = client.get(f"/api/v1/capex/cost-centers/{cc['id']}", headers=admin).json()
    view = client.get(f"/api/v1/capex/submissions/{capex_head['submission_id']}/view", headers=admin).json()
    assert view["totals"]["proposed"] == "50000.00"
    ov11 = client.get("/api/v1/consolidation/overview", headers=admin).json()
    assert ov11["version"]["label"] == "1.1"
    # só ANA ficou ativa no quadro (226.800) + a vaga, que foi copiada da 1.0 para a revisão (34.020)
    assert ov11["modules"]["PERSONNEL"]["proposed"] == "260820.00"
    v10 = next(v for v in ov11["versions"] if v["label"] == "1.0")
    old = client.get(f"/api/v1/consolidation/overview?version_id={v10['id']}", headers=admin).json()
    assert old["total"] == "524750.00" and old["version"]["status"] == "FROZEN"


def test_manager_sees_only_own_cost_centers(client, admin, run_worker):
    from tests.test_opex import _user

    cc, _, _ = _setup(client, admin, run_worker)
    _, other = _user(client, admin, "outro@t.com", ["MANAGER"])
    ov = client.get("/api/v1/consolidation/overview", headers=other).json()
    assert ov["total"] == "0.00" and ov["matrix"] == []
    wb = load_workbook(io.BytesIO(client.get("/api/v1/consolidation/export.xlsx", headers=other).content))
    assert wb["Carga SAP"].max_row == 1


def test_delete_all_resets_versions(client, admin, run_worker):
    _setup(client, admin, run_worker)
    client.post("/api/v1/consolidation/freeze", headers=admin, json={})
    client.post("/api/v1/consolidation/revise", headers=admin, json={"reason": "teste"})
    r = client.delete("/api/v1/datasets/all", headers=admin, params={"confirm": "EXCLUIR"})
    assert r.status_code == 200, r.text
    ov = client.get("/api/v1/consolidation/overview", headers=admin).json()
    assert [(v["label"], v["status"]) for v in ov["versions"]] == [("1.0", "WORKING")]
    assert ov["total"] == "0.00"


def test_frozen_version_does_not_create_submissions(client, admin, run_worker):
    _setup(client, admin, run_worker)
    client.post("/api/v1/consolidation/freeze", headers=admin, json={})
    other = next(c for c in client.get("/api/v1/cost-centers", headers=admin).json() if c["code"] == "1050101012")
    for module in ("opex", "capex", "personnel"):
        resp = client.get(f"/api/v1/{module}/cost-centers/{other['id']}", headers=admin)
        assert resp.status_code == 404 and "congelada" in resp.json()["detail"], module
    # nenhum orçamento novo apareceu na versão congelada
    ov = client.get("/api/v1/consolidation/overview", headers=admin).json()
    row = next(r for r in ov["matrix"] if r["code"] == "1050101012")
    assert row["status"] == {"OPEX": "DRAFT", "CAPEX": "DRAFT", "PERSONNEL": "DRAFT"}
    # excluir orçamento por módulo é recusado com a versão congelada; excluir tudo zera e volta à 1.0
    r = client.delete("/api/v1/datasets/budget", headers=admin, params={"module": "OPEX", "confirm": "EXCLUIR"})
    assert r.status_code == 409
    assert client.delete("/api/v1/datasets/all", headers=admin, params={"confirm": "EXCLUIR"}).status_code == 200


def test_delete_all_with_two_revisions(client, admin, run_worker):
    _setup(client, admin, run_worker)
    for n in (1, 2):
        client.post("/api/v1/consolidation/freeze", headers=admin, json={})
        assert client.post("/api/v1/consolidation/revise", headers=admin, json={"reason": f"r{n}"}).status_code == 200
    labels = [v["label"] for v in client.get("/api/v1/consolidation/overview", headers=admin).json()["versions"]]
    assert labels == ["1.0", "1.1", "1.2"]
    r = client.delete("/api/v1/datasets/all", headers=admin, params={"confirm": "EXCLUIR"})
    assert r.status_code == 200, r.text
    assert [v["label"] for v in client.get("/api/v1/consolidation/overview", headers=admin).json()["versions"]] == [
        "1.0"
    ]


def test_compare_versions(client, admin, run_worker):
    cc, _, line = _setup(client, admin, run_worker)
    client.post("/api/v1/consolidation/freeze", headers=admin, json={})
    rv = client.post("/api/v1/consolidation/revise", headers=admin, json={"reason": "ajuste"}).json()
    v10, v11 = (next(v["id"] for v in rv["versions"] if v["label"] == lb) for lb in ("1.0", "1.1"))
    head = client.get(f"/api/v1/opex/cost-centers/{cc['id']}", headers=admin).json()
    new_line = client.get(f"/api/v1/opex/submissions/{head['submission_id']}/lines", headers=admin).json()[0]
    client.patch(f"/api/v1/opex/lines/{new_line['id']}", headers=admin, json={"values": {"1": 1450}})  # +1.000 em JAN
    r = client.get(f"/api/v1/consolidation/compare?from_version_id={v10}&to_version_id={v11}", headers=admin).json()
    assert r["difference"] == "1000.00" and r["monthly_difference"][0] == "1000.00"
    assert r["modules"] == [
        {
            "module": "OPEX",
            "label": "OPEX",
            "from": "5400.00",
            "to": "6400.00",
            "difference": "1000.00",
            "difference_pct": "0.1852",
        }
    ]
    assert r["accounts"][0]["account"] == "6010301002" and r["changed_accounts"] == 1
    assert (
        client.get(
            f"/api/v1/consolidation/compare?from_version_id={v10}&to_version_id={v10}", headers=admin
        ).status_code
        == 422
    )
