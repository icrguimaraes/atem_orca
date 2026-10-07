from tests import builders
from tests.test_imports import import_and_load, status, upload


def test_filled_opex_template_loads_everything(client, admin, run_worker):
    batch_id = upload(client, admin, builders.opex_template_filled(), "Template_OPEX 2027_Gestor.xlsx")
    run_worker()
    b = status(client, admin, batch_id)
    assert b["dataset_type"] == "OPEX_TEMPLATE", b
    meta = b["summary"]["meta"]
    assert meta["parts"] == {"master": 9, "actual": 1, "budget_lines": 4}
    budget = b["summary"]["comparison"]["budget"]
    assert budget == [
        {
            "cost_center": "1050101011",
            "company": "1001",
            "status": "DRAFT",
            "editable": True,
            "lines": 4,
            "total": "20400",
            "replaces_lines": 0,
            "replaces_total": "0",
        }
    ]
    assert client.post(f"/api/v1/imports/{batch_id}/confirm", headers=admin).status_code == 200
    run_worker()
    final = status(client, admin, batch_id)
    assert final["status"] == "COMPLETED", final
    assert final["summary"]["load"]["budget"]["lines_created"] == 4

    cc = next(c for c in client.get("/api/v1/cost-centers", headers=admin).json() if c["code"] == "1050101011")
    head = client.get(f"/api/v1/opex/cost-centers/{cc['id']}", headers=admin).json()
    assert head["status"] == "IN_PROGRESS"
    view = client.get(f"/api/v1/opex/submissions/{head['submission_id']}/accounts", headers=admin).json()
    rows = {r["code"]: r for r in view["accounts"]}
    assert rows["6010301002"]["proposed"] == "4200.00" and rows["6010301002"]["ref_actual_ytd"] == "2400.00"
    assert rows["6010301011"]["proposed"] == "1800.00"
    assert meta["consolidator_check"]["I - Viagens"]["ok"] is True
    lines = client.get(f"/api/v1/opex/submissions/{head['submission_id']}/lines", headers=admin).json()
    trips = [line for line in lines if line["line_type"] == "TRAVEL"]
    assert {t["description"] for t in trips} == {"Auditoria SP", "Visita base Belém"}
    belem = next(t for t in trips if t["description"] == "Visita base Belém")
    assert belem["attributes"]["destination"] == "PA" and belem["attributes"]["return_month"] is None
    assert belem["values"]["3"] == "2400.00"
    # passagem zerada entre lugares diferentes: apontada na viagem (tela) e na prévia da importação
    assert belem["attributes"]["warnings"] == ["Passagem orçada em R$ 0: a planilha não trouxe tarifa para AM → PA"]
    auditoria = next(t for t in trips if t["description"] == "Auditoria SP")
    assert auditoria["attributes"]["warnings"] == []
    errors = client.get(f"/api/v1/imports/{batch_id}/errors", headers=admin).json()
    codes = {e["code"] for e in (errors["items"] if isinstance(errors, dict) else errors)}
    assert "TRAVEL_NO_FARE" in codes
    consult = next(line for line in lines if line["supplier"] == "KPMG")
    assert consult["description"] == "Consultoria tributária" and consult["justification"] == "Reajuste IPCA"

    # linha digitada no sistema é preservada; reimportar o template substitui só as linhas do template
    acc = next(
        a for a in client.get("/api/v1/opex/options", headers=admin).json()["accounts"] if a["code"] == "6010301005"
    )
    client.post(
        f"/api/v1/opex/submissions/{head['submission_id']}/lines",
        headers=admin,
        json={"account_id": acc["id"], "values": {"1": 100}},
    )
    final = import_and_load(client, admin, run_worker, builders.opex_template_filled(), "v2.xlsx", force=True)
    assert final["summary"]["load"]["budget"] == {
        "lines_replaced": 4,
        "lines_created": 4,
        "cost_centers": 1,
        "trips": 2,
    }
    lines = client.get(f"/api/v1/opex/submissions/{head['submission_id']}/lines", headers=admin).json()
    assert len(lines) == 5


def test_template_blocked_after_submission(client, admin, run_worker):
    import_and_load(client, admin, run_worker, builders.opex_template_filled(), "t.xlsx")
    cc = next(c for c in client.get("/api/v1/cost-centers", headers=admin).json() if c["code"] == "1050101011")
    head = client.get(f"/api/v1/opex/cost-centers/{cc['id']}", headers=admin).json()
    sub = head["submission_id"]
    view = client.get(f"/api/v1/opex/submissions/{sub}/accounts", headers=admin).json()
    for r in view["accounts"]:
        if r["needs_justification"]:
            client.put(
                f"/api/v1/opex/submissions/{sub}/justifications/{r['account_id']}", headers=admin, json={"text": "ok"}
            )
    assert client.post(f"/api/v1/opex/submissions/{sub}/actions/submit", headers=admin, json={}).status_code == 200
    batch_id = upload(client, admin, builders.opex_template_filled(), "t2.xlsx")
    run_worker()
    errors = client.get(f"/api/v1/imports/{batch_id}/errors", headers=admin, params={"code": "BUDGET_LOCKED"}).json()
    assert len(errors) == 4


def test_broken_key_resolved_by_name_and_unresolved_is_reported(client, admin, run_worker):
    """CHAVE com CC = 0 (XLOOKUP não achou): identifica pelo nome; sem nome válido, erro visível."""
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(builders.opex_template_filled()))
    ws = wb["II - Serviços de Terceiros"]
    row = next(r for r in range(15, ws.max_row + 1) if ws.cell(r, 5).value == "KPMG")
    ws.cell(row, 2, "1001-0001-0-6010201003")
    ws.cell(row, 10, 0)
    ws.cell(row, 9, "DADOS E PROJ. APLICADOS A CONTROLADORIA")
    ws.append(
        [
            None,
            "1001-0001-0-6010201003",
            "Linha órfã",
            None,
            None,
            None,
            "MANAUS",
            "0001",
            "CC INEXISTENTE",
            0,
            "Serviços de Consultoria e Assessoria",
            "6010201003",
            *[500] * 12,
            6000,
        ]
    )
    via = wb["I - Viagens"]
    via.cell(63, 19, 2000)  # hospedagem alterada sem atualizar o consolidador
    buf = io.BytesIO()
    wb.save(buf)

    batch_id = upload(client, admin, buf.getvalue(), "t.xlsx")
    run_worker()
    b = status(client, admin, batch_id)
    assert b["summary"]["comparison"]["budget"][0]["lines"] == 4
    errors = client.get(f"/api/v1/imports/{batch_id}/errors", headers=admin).json()
    items = errors["items"] if isinstance(errors, dict) else errors
    codes = {e["code"] for e in items}
    assert {"UNRESOLVED_LINE", "CONSOLIDATOR_MISMATCH"} <= codes
    assert b["summary"]["meta"]["consolidator_check"]["I - Viagens"]["ok"] is False
